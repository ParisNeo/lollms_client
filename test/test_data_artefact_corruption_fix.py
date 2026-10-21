import unittest
import sys
import os
import tempfile
import shutil
import sqlite3
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from lollms_client.lollms_discussion import LollmsDiscussion, LollmsDataManager
from lollms_client.lollms_artefact import ArtefactType
from lollms_client.tools_bindings.lcp.default_tools.execute_python_data_query.execute_python_data_query import (
    tool_execute_python_data_query,
    _validate_xlsx,
    _validate_office_file,
)


class MockClient:
    """Mock client for testing artifact operations without an LLM."""
    def __init__(self):
        self.llm = self
        self.ai_name = "Assistant"
        self.model_name = "mock"
        self.binding_name = "mock"

    def count_tokens(self, text: str) -> int:
        return len(text) // 4

    def count_image_tokens(self, img) -> int:
        return 0

    def remove_thinking_blocks(self, text: str) -> str:
        return text

    def generate_text(self, prompt: str, **kwargs) -> str:
        return "ok"

    def reset_cancel(self):
        pass


class TestDataArtefactCorruptionFix(unittest.TestCase):
    """
    Validates that updating a DATA artifact (e.g., SQLite .db) does not corrupt
    the physical binary file by overwriting it with the .lam logical schema text.
    """

    def setUp(self):
        self.client = MockClient()
        self.db_manager = LollmsDataManager("sqlite:///:memory:")
        self.tmp_dir = tempfile.mkdtemp(prefix="lollms_data_corruption_")
        self.discussion = LollmsDiscussion.create_new(
            lollms_client=self.client,
            db_manager=self.db_manager,
            id="data_corruption_test",
            workspace_path=self.tmp_dir,
            autosave=True
        )

        # Create a valid SQLite database in a temporary file and read its bytes
        self.tmp_db_path = Path(self.tmp_dir) / "source_test.db"
        self.conn = sqlite3.connect(str(self.tmp_db_path))
        self.conn.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, name TEXT)")
        self.conn.execute("INSERT INTO users (name) VALUES ('Alice')")
        self.conn.commit()
        self.conn.close()
        self.db_bytes = self.tmp_db_path.read_bytes()

        # The expected valid SQLite header
        self.valid_header = b"SQLite format 3\x00"

    def tearDown(self):
        self.discussion.close()
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def test_update_preserves_sqlite_binary_integrity(self):
        """
        Scenario: A SQLite database is imported as a DATA artifact. The LLM updates
        the artifact's logical content (.lam schema).
        Expected: The physical .db file on disk must remain a valid SQLite database.
                  It must NOT be overwritten with the .lam schema text.
        """
        title = "test_db.db"
        lam_content = "# SQLite Database: test_db.db\n## Table: users\nColumns: id (INTEGER), name (TEXT)"

        # 1. Add the data artifact with physical bytes
        self.discussion.artefacts.add(
            title=title,
            artefact_type=ArtefactType.DATA,
            content=lam_content,
            logical_content=lam_content,
            physical_data=self.db_bytes,
            file_ext=".db",
            active=True
        )
        self.discussion.commit()

        # Verify initial state is valid
        ws_data_dir = Path(self.discussion.workspace_data_path)
        db_path = ws_data_dir / title
        self.assertTrue(db_path.exists(), "Physical .db file should exist after add()")
        self.assertEqual(db_path.read_bytes()[:16], self.valid_header, "Initial DB should have valid SQLite header")

        # 2. Update the artifact's logical content (simulates LLM editing the .lam)
        updated_lam = "# SQLite Database: test_db.db\n## Table: users\nColumns: id (INTEGER), name (TEXT), email (TEXT)"
        self.discussion.artefacts.update(
            title=title,
            new_content=updated_lam,
            logical_content=updated_lam,
            bump_version=True
        )
        self.discussion.commit()

        # 3. CRITICAL ASSERTION: The physical file must still be a valid SQLite database
        self.assertTrue(db_path.exists(), "Physical .db file should still exist after update()")
        header = db_path.read_bytes()[:16]
        self.assertEqual(header, self.valid_header,
                         f"Database corrupted after update! Header is '{header}' instead of '{self.valid_header}'. "
                         f"The .lam schema text was likely written to the .db file.")

        # 4. Verify we can still query the database
        try:
            conn = sqlite3.connect(str(db_path))
            cursor = conn.cursor()
            cursor.execute("SELECT name FROM users WHERE id = 1")
            row = cursor.fetchone()
            self.assertIsNotNone(row, "Database query should return a row")
            self.assertEqual(row[0], "Alice", "Database content should be preserved")
            conn.close()
        except sqlite3.DatabaseError as e:
            self.fail(f"Failed to query database after update: {e}. File may be corrupted.")

    def test_update_without_logical_content_preserves_sqlite_binary(self):
        """
        Scenario: A SQLite database is imported as a DATA artifact. The user edits
        the artifact's metadata via the UI Raw Editor, which calls update() with
        new_content = edited schema text but does NOT pass logical_content.
        Expected: The is_raw_data_update heuristic must NOT treat the string
                  new_content as raw binary data for .db files. The physical .db
                  file on disk must remain a valid SQLite database.
        """
        title = "test_db_no_lam.db"
        lam_content = "# SQLite Database: test_db_no_lam.db\n## Table: users\nColumns: id (INTEGER), name (TEXT)"

        # 1. Add the data artifact with physical bytes
        self.discussion.artefacts.add(
            title=title,
            artefact_type=ArtefactType.DATA,
            content=lam_content,
            logical_content=lam_content,
            physical_data=self.db_bytes,
            file_ext=".db",
            active=True
        )
        self.discussion.commit()

        # Verify initial state is valid
        ws_data_dir = Path(self.discussion.workspace_data_path)
        db_path = ws_data_dir / title
        self.assertTrue(db_path.exists(), "Physical .db file should exist after add()")
        self.assertEqual(db_path.read_bytes()[:16], self.valid_header, "Initial DB should have valid SQLite header")

        # 2. Update the artifact's content WITHOUT passing logical_content.
        #    This simulates the UI Raw Editor path (POST /api/artifacts/{title}/update)
        #    which passes new_content=payload.content but no logical_content parameter.
        #    The is_raw_data_update heuristic must detect _is_binary_db and refuse
        #    to encode the schema text as physical bytes.
        updated_lam = "# SQLite Database: test_db_no_lam.db\n## Table: users\nColumns: id (INTEGER), name (TEXT), email (TEXT)"
        self.discussion.artefacts.update(
            title=title,
            new_content=updated_lam,
            bump_version=True
        )
        self.discussion.commit()

        # 3. CRITICAL ASSERTION: The physical file must still be a valid SQLite database
        self.assertTrue(db_path.exists(), "Physical .db file should still exist after update()")
        header = db_path.read_bytes()[:16]
        self.assertEqual(header, self.valid_header,
                         f"Database corrupted after update without logical_content! "
                         f"Header is '{header}' instead of '{self.valid_header}'. "
                         f"The is_raw_data_update heuristic encoded the .lam schema text as binary bytes.")

        # 4. Verify we can still query the database
        try:
            conn = sqlite3.connect(str(db_path))
            cursor = conn.cursor()
            cursor.execute("SELECT name FROM users WHERE id = 1")
            row = cursor.fetchone()
            self.assertIsNotNone(row, "Database query should return a row")
            self.assertEqual(row[0], "Alice", "Database content should be preserved")
            conn.close()
        except sqlite3.DatabaseError as e:
            self.fail(f"Failed to query database after update without logical_content: {e}. File may be corrupted.")

        # 5. Verify the artifact's content field was updated to the new schema
        updated_art = self.discussion.artefacts.get(title)
        self.assertIsNotNone(updated_art)
        self.assertIn("email", updated_art.get("content", ""),
                      "The artifact content should reflect the updated schema with the 'email' column.")

    def test_sync_to_disk_refuses_text_for_binary_db(self):
        """
        Scenario: _sync_to_disk_workspace is called directly with string content
        but no physical_data for a .db file.
        Expected: The method must refuse to write the string content to the .db file
                  to prevent silent corruption.
        """
        title = "protected_db.db"
        lam_content = "# This is schema text, not binary data"

        # Call the internal sync method directly with no physical_data
        self.discussion.artefacts._sync_to_disk_workspace(
            title=title,
            content=lam_content,
            version=1,
            atype=ArtefactType.DATA,
            file_ext=".db",
            physical_data=None,
            logical_content=lam_content
        )

        ws_data_dir = Path(self.discussion.workspace_data_path)
        db_path = ws_data_dir / title

        # The .db file should either NOT exist, or if it does, it must NOT contain the schema text
        if db_path.exists():
            content = db_path.read_bytes()
            self.assertNotIn(lam_content.encode('utf-8'), content,
                             "Schema text was written to the .db file! Binary corruption occurred.")
        # If it doesn't exist, the defense-in-depth worked by refusing to write.

    def test_update_preserves_csv_physical_bytes(self):
        """
        Scenario: A CSV file is imported as a DATA artifact. The LLM updates the schema.
        Expected: The physical .csv file on disk must contain the original CSV rows,
                  not the .lam schema text.
        """
        title = "data.csv"
        csv_bytes = b"id,name\n1,Alice\n2,Bob\n"
        lam_content = "# CSV Schema: data.csv\nColumns: id (int), name (str)"

        # 1. Add the data artifact
        self.discussion.artefacts.add(
            title=title,
            artefact_type=ArtefactType.DATA,
            content=lam_content,
            logical_content=lam_content,
            physical_data=csv_bytes,
            file_ext=".csv",
            active=True
        )
        self.discussion.commit()

        # 2. Update the logical schema
        updated_lam = "# CSV Schema: data.csv\nColumns: id (int), name (str), email (str)"
        self.discussion.artefacts.update(
            title=title,
            new_content=updated_lam,
            logical_content=updated_lam,
            bump_version=True
        )
        self.discussion.commit()

        # 3. Verify physical file still contains CSV data
        ws_data_dir = Path(self.discussion.workspace_data_path)
        csv_path = ws_data_dir / title
        self.assertTrue(csv_path.exists(), "Physical .csv file should exist after update()")

        content = csv_path.read_bytes()
        self.assertEqual(content, csv_bytes,
                         "Physical CSV bytes should be exactly preserved after update().")
        self.assertNotIn(b"# CSV Schema", content,
                         "The .lam schema text was written to the physical .csv file! Corruption detected.")


class TestUpdatePhantomRevertFix(unittest.TestCase):
    """
    Validates that ArtefactManager.update() does not create a phantom version 
    with old content after updating with new content (simulates the FastAPI 
    update_manual_artefact endpoint behavior).
    """

    def setUp(self):
        self.client = MockClient()
        self.db_manager = LollmsDataManager("sqlite:///:memory:")
        import tempfile
        self.tmp_dir = tempfile.mkdtemp(prefix="lollms_phantom_revert_")
        self.discussion = LollmsDiscussion.create_new(
            lollms_client=self.client,
            db_manager=self.db_manager,
            id="phantom_revert_test",
            workspace_path=self.tmp_dir,
            autosave=True
        )

    def tearDown(self):
        self.discussion.close()
        import shutil
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def test_update_does_not_create_phantom_version(self):
        """
        Scenario: An artifact is added (v1), then updated with new content using 
                  bump_version=False and active=True (mimicking the router call).
        Expected: Only 1 version exists in the database, containing the NEW content.
                  The physical file on disk also contains the NEW content.
        """
        title = "config.txt"
        original_content = "initial_value = 1"
        updated_content = "initial_value = 2"

        # 1. Add the artifact (creates v1)
        self.discussion.artefacts.add(
            title=title,
            artefact_type=ArtefactType.DOCUMENT,
            content=original_content,
            active=True
        )
        self.discussion.commit()

        # Verify initial state
        history = self.discussion.artefacts.get_version_history(title)
        self.assertEqual(len(history), 1, "Precondition: Should have 1 version after add()")

        # 2. Update the artifact (simulates router call: bump_version=False, active=True)
        self.discussion.artefacts.update(
            title=title,
            new_content=updated_content,
            bump_version=False,
            active=True
        )
        self.discussion.commit()

        # 3. Assert NO phantom version was created
        history_after = self.discussion.artefacts.get_version_history(title)
        self.assertEqual(len(history_after), 1, 
                         "Phantom version created! update() should not add a new version when bump_version=False.")

        # 4. Assert the DB content is the NEW content
        latest_art = self.discussion.artefacts.get(title)
        self.assertIsNotNone(latest_art)
        self.assertEqual(latest_art["content"], updated_content,
                         "DB content should be the updated content, not the original.")

        # 5. Assert the physical file on disk contains the NEW content
        ws_data_dir = Path(self.discussion.workspace_data_path)
        file_path = ws_data_dir / title
        self.assertTrue(file_path.exists(), "Physical file should exist after update.")
        
        disk_content = file_path.read_text(encoding="utf-8")
        self.assertEqual(disk_content, updated_content,
                         "Physical file on disk should contain the updated content, not the original.")
        self.assertNotIn(original_content, disk_content,
                         "Physical file on disk should NOT contain the original content.")


class TestExecutePythonDataQueryXlsxGuard(unittest.TestCase):
    """
    Validates the transactional XLSX guard in execute_python_data_query:
    snapshot before execution, strict three-layer validation after execution,
    and byte-identical rollback (or deletion of corrupt new files) so no
    Excel-unreadable spreadsheet is ever committed to the workspace.
    """

    def setUp(self):
        self.original_cwd = os.getcwd()
        self.tmp_dir = tempfile.mkdtemp(prefix="lollms_xlsx_guard_")
        os.chdir(self.tmp_dir)

        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Feuil1"
        ws["A1"] = "original"
        wb.save("to_fill.xlsx")
        self.original_bytes = Path("to_fill.xlsx").read_bytes()

    def tearDown(self):
        os.chdir(self.original_cwd)
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def test_valid_mutation_passes_guard(self):
        """
        Scenario: user code loads the workbook, adds a sheet, saves.
        Expected: success, file remains openable, no rollback occurred.
        """
        result = tool_execute_python_data_query(
            "import openpyxl\n"
            "wb = openpyxl.load_workbook('to_fill.xlsx')\n"
            "wb.create_sheet('games')\n"
            "wb.save('to_fill.xlsx')\n"
            "print('saved')"
        )
        self.assertTrue(result["success"], f"Valid mutation should succeed: {result.get('error')}")
        self.assertNotIn("xlsx_failed_excel_validation", result.get("output", ""))

        import openpyxl
        wb = openpyxl.load_workbook("to_fill.xlsx")
        self.assertIn("games", wb.sheetnames)

    def test_corrupt_mutation_is_rolled_back(self):
        """
        Scenario: user code overwrites the workbook with garbage bytes.
        Expected: structured failure with the xlsx_failed_excel_validation code,
                  and the file restored byte-identically to its pre-execution state.
        """
        result = tool_execute_python_data_query(
            "with open('to_fill.xlsx', 'wb') as f:\n"
            "    f.write(b'PK\\x03\\x04garbage-not-a-zip')\n"
            "print('overwritten')"
        )
        self.assertFalse(result["success"], "Corrupting code must not report success.")
        self.assertIn("xlsx_failed_excel_validation", result["error"])
        self.assertIn("rolled back", result["error"].lower())

        restored = Path("to_fill.xlsx").read_bytes()
        self.assertEqual(restored, self.original_bytes,
                         "File must be restored byte-identically to the pre-execution snapshot.")

    def test_legitimate_pie_chart_passes_guard(self):
        """
        Scenario: user code adds a PieChart with geometry set BEFORE add_chart,
                  anchored at a free cell away from data ranges.
        Expected: success — the guard must not over-block legitimate chart workloads.
        """
        result = tool_execute_python_data_query(
            "import openpyxl\n"
            "from openpyxl.chart import PieChart, Reference\n"
            "wb = openpyxl.load_workbook('to_fill.xlsx')\n"
            "ws = wb.create_sheet('charts')\n"
            "ws['A1'] = 'Genre'\n"
            "ws['B1'] = 'Hours'\n"
            "ws['A2'] = 'RPG'\n"
            "ws['B2'] = 120\n"
            "ws['A3'] = 'FPS'\n"
            "ws['B3'] = 25\n"
            "pie = PieChart()\n"
            "pie.title = 'Playtime by Genre'\n"
            "pie.width = 15\n"
            "pie.height = 10\n"
            "data = Reference(ws, min_col=2, min_row=1, max_row=3)\n"
            "cats = Reference(ws, min_col=1, min_row=2, max_row=3)\n"
            "pie.add_data(data, titles_from_data=True)\n"
            "pie.set_categories(cats)\n"
            "ws.add_chart(pie, 'D2')\n"
            "wb.save('to_fill.xlsx')\n"
            "print('chart saved')"
        )
        self.assertTrue(result["success"], f"Legitimate chart should pass the guard: {result.get('error')}")
        ok, reason = _validate_xlsx(Path("to_fill.xlsx"))
        self.assertTrue(ok, f"Chart workbook should validate: {reason}")

    def test_corrupt_new_file_is_deleted(self):
        """
        Scenario: user code creates a brand-new corrupt .xlsx (no snapshot exists).
        Expected: structured failure and the corrupt file deleted from the workspace.
        """
        result = tool_execute_python_data_query(
            "with open('new_bad.xlsx', 'wb') as f:\n"
            "    f.write(b'PK\\x03\\x04garbage-not-a-zip')\n"
            "print('created')"
        )
        self.assertFalse(result["success"], "Creating a corrupt new file must fail.")
        self.assertIn("xlsx_failed_excel_validation", result["error"])
        self.assertFalse(Path("new_bad.xlsx").exists(),
                         "Corrupt new file must be deleted, not committed to the workspace.")

    def test_auto_df_excel_persistence_is_validated(self):
        """
        Scenario: code references df without defining it (auto-load triggers on the
                  .xlsx in CWD), mutates it, and the tool persists via to_excel.
        Expected: success and the persisted file passes validation.
        """
        import pandas as pd
        source_df = pd.DataFrame({"name": ["Alice", "Bob"], "score": [10, 20]})
        source_df.to_excel("data.xlsx", index=False)

        result = tool_execute_python_data_query(
            "df['score'] = df['score'] * 2\n"
            "print(df.head())"
        )
        self.assertTrue(result["success"], f"Auto-df persistence should succeed: {result.get('error')}")
        ok, reason = _validate_xlsx(Path("data.xlsx"))
        self.assertTrue(ok, f"Persisted df workbook should validate: {reason}")


class TestExecutePythonDataQueryOfficeGuard(unittest.TestCase):
    """
    Validates the generalized transactional Office guard in execute_python_data_query:
    snapshot before execution, strict three-layer validation after execution, and
    byte-identical rollback (or deletion of corrupt new files) for .docx/.pptx/.xlsx
    so no Word/Excel/PowerPoint-unreadable document is ever committed to the workspace.
    """

    def setUp(self):
        self.original_cwd = os.getcwd()
        self.tmp_dir = tempfile.mkdtemp(prefix="lollms_office_guard_")
        os.chdir(self.tmp_dir)

        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Feuil1"
        ws["A1"] = "original"
        wb.save("to_fill.xlsx")
        self.original_xlsx_bytes = Path("to_fill.xlsx").read_bytes()

        try:
            from docx import Document
            doc = Document()
            doc.add_paragraph("original")
            doc.save("to_fill.docx")
            self.original_docx_bytes = Path("to_fill.docx").read_bytes()
        except ImportError:
            self.original_docx_bytes = None

        try:
            from pptx import Presentation
            prs = Presentation()
            prs.slide_layouts[6]
            prs.slides.add_slide(prs.slide_layouts[6])
            prs.save("to_fill.pptx")
            self.original_pptx_bytes = Path("to_fill.pptx").read_bytes()
        except ImportError:
            self.original_pptx_bytes = None

    def tearDown(self):
        os.chdir(self.original_cwd)
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def test_valid_docx_mutation_passes_guard(self):
        """
        Scenario: user code loads the document, adds a paragraph, saves.
        Expected: success, file remains openable, no rollback occurred.
        """
        if self.original_docx_bytes is None:
            self.skipTest("python-docx not installed")
        result = tool_execute_python_data_query(
            "from docx import Document\n"
            "doc = Document('to_fill.docx')\n"
            "doc.add_paragraph('added by guard test')\n"
            "doc.save('to_fill.docx')\n"
            "print('saved')"
        )
        self.assertTrue(result["success"], f"Valid docx mutation should succeed: {result.get('error')}")
        ok, reason = _validate_office_file(Path("to_fill.docx"))
        self.assertTrue(ok, f"Mutated docx should validate: {reason}")

    def test_valid_pptx_mutation_passes_guard(self):
        """
        Scenario: user code loads the presentation, adds a slide, saves.
        Expected: success, file remains openable, no rollback occurred.
        """
        if self.original_pptx_bytes is None:
            self.skipTest("python-pptx not installed")
        result = tool_execute_python_data_query(
            "from pptx import Presentation\n"
            "prs = Presentation('to_fill.pptx')\n"
            "prs.slides.add_slide(prs.slide_layouts[6])\n"
            "prs.save('to_fill.pptx')\n"
            "print('saved')"
        )
        self.assertTrue(result["success"], f"Valid pptx mutation should succeed: {result.get('error')}")
        ok, reason = _validate_office_file(Path("to_fill.pptx"))
        self.assertTrue(ok, f"Mutated pptx should validate: {reason}")

    def test_corrupt_docx_mutation_is_rolled_back(self):
        """
        Scenario: user code overwrites the document with garbage bytes.
        Expected: structured failure with the docx_failed_word_validation code,
                  and the file restored byte-identically to its pre-execution state.
        """
        if self.original_docx_bytes is None:
            self.skipTest("python-docx not installed")
        result = tool_execute_python_data_query(
            "with open('to_fill.docx', 'wb') as f:\n"
            "    f.write(b'PK\\x03\\x04garbage-not-a-zip')\n"
            "print('overwritten')"
        )
        self.assertFalse(result["success"], "Corrupting code must not report success.")
        self.assertIn("docx_failed_word_validation", result["error"])
        self.assertIn("rolled back", result["error"].lower())

        restored = Path("to_fill.docx").read_bytes()
        self.assertEqual(restored, self.original_docx_bytes,
                         "File must be restored byte-identically to the pre-execution snapshot.")

    def test_corrupt_pptx_mutation_is_rolled_back(self):
        """
        Scenario: user code overwrites the presentation with garbage bytes.
        Expected: structured failure with the pptx_failed_powerpoint_validation code,
                  and the file restored byte-identically to its pre-execution state.
        """
        if self.original_pptx_bytes is None:
            self.skipTest("python-pptx not installed")
        result = tool_execute_python_data_query(
            "with open('to_fill.pptx', 'wb') as f:\n"
            "    f.write(b'PK\\x03\\x04garbage-not-a-zip')\n"
            "print('overwritten')"
        )
        self.assertFalse(result["success"], "Corrupting code must not report success.")
        self.assertIn("pptx_failed_powerpoint_validation", result["error"])
        self.assertIn("rolled back", result["error"].lower())

        restored = Path("to_fill.pptx").read_bytes()
        self.assertEqual(restored, self.original_pptx_bytes,
                         "File must be restored byte-identically to the pre-execution snapshot.")

    def test_corrupt_new_pptx_file_is_deleted(self):
        """
        Scenario: user code creates a brand-new corrupt .pptx (no snapshot exists).
        Expected: structured failure and the corrupt file deleted from the workspace.
        """
        result = tool_execute_python_data_query(
            "with open('new_bad.pptx', 'wb') as f:\n"
            "    f.write(b'PK\\x03\\x04garbage-not-a-zip')\n"
            "print('created')"
        )
        self.assertFalse(result["success"], "Creating a corrupt new file must fail.")
        self.assertIn("pptx_failed_powerpoint_validation", result["error"])
        self.assertFalse(Path("new_bad.pptx").exists(),
                         "Corrupt new file must be deleted, not committed to the workspace.")

    def test_xlsx_regression_intact_under_generalized_guard(self):
        """
        Scenario: the original xlsx corrupt-rollback scenario re-run under the
                  generalized Office guard.
        Expected: identical behavior to the dedicated xlsx suite — failure with
                  xlsx_failed_excel_validation and byte-identical restoration.
        """
        result = tool_execute_python_data_query(
            "with open('to_fill.xlsx', 'wb') as f:\n"
            "    f.write(b'PK\\x03\\x04garbage-not-a-zip')\n"
            "print('overwritten')"
        )
        self.assertFalse(result["success"], "Corrupting code must not report success.")
        self.assertIn("xlsx_failed_excel_validation", result["error"])

        restored = Path("to_fill.xlsx").read_bytes()
        self.assertEqual(restored, self.original_xlsx_bytes,
                         "File must be restored byte-identically to the pre-execution snapshot.")


if __name__ == "__main__":
    unittest.main()
