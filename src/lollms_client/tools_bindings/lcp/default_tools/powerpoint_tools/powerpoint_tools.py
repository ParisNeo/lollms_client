"""
PowerPoint Tools — dedicated LCP toolset for Microsoft PowerPoint decks.

Doctrine (read this before writing any recipe):
- SLIDE SIZE: the pre-bound root presentation is 16:9 — 13.333in x 7.5in
  (Emu 12192000 x 6858000). Position shapes within these bounds.
- TEMPLATE AWARENESS: call tool_pptx_info to discover the deck's slide
  layouts (names + placeholder types) and theme fonts before building.
- VISUAL VERIFICATION: call tool_view_slides (max 4 slides per call) to SEE
  rendered slides. The images are attached to the CURRENT ROUND ONLY — they
  are never stored in the workspace.
"""
import base64
import re
import shutil
import subprocess
import tempfile
import textwrap
import traceback
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from ascii_colors import ASCIIColors

TOOL_LIBRARY_NAME = "PowerPoint Tools"
TOOL_LIBRARY_DESC = (
    "Create, edit, read and visually inspect PowerPoint decks (.pptx/.pptm) via "
    "sandboxed python-pptx recipe code. The root presentation is 16:9 (13.333in x "
    "7.5in). Use tool_pptx_info for template/layout awareness and tool_view_slides "
    "(max 4 slides per call, transient images) for visual verification. Every save "
    "is verified with a strict three-layer OOXML integrity check."
)
TOOL_LIBRARY_ICON = "📊"

_MAX_RECIPE_CHARS = 60_000
_FORBIDDEN_TOKENS = (
    "eval(", "exec(", "compile(", "open(",
    "os.", "sys.", "subprocess", "shutil", "pathlib", "__",
    "globals(", "locals(", "vars(", "getattr(", "setattr(", "delattr(",
)

_IMPORT_PREFIX_WHITELIST = ("pptx.",)
_SAFE_STDLIB_MODULES = frozenset({
    "math", "datetime", "random", "json", "statistics", "itertools",
    "functools", "re", "uuid", "calendar", "decimal", "fractions",
    "textwrap", "base64", "hashlib", "string", "collections",
})

_MEDIA_EXTENSIONS = (
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg", ".tiff", ".ico",
    ".mp4", ".mov", ".avi", ".mkv", ".webm", ".wmv",
)
_MAX_MEDIA_SCAN = 5000
_MAX_MEDIA_LIST = 2000

OOXML_VALIDATORS = {
    ".pptx": {"required_parts": ("[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"), "reopener": "python-pptx"},
    ".pptm": {"required_parts": ("[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"), "reopener": "python-pptx"},
}
GUARDED_EXTENSIONS = tuple(OOXML_VALIDATORS.keys())

OFFICE_GUARD_FAILURE_CODE = "office_file_failed_validation"

OFFICE_GUARD_GUIDANCE = (
    "Likely causes: python-pptx re-saves dropping parts desktop PowerPoint requires "
    "(slide layouts, master rels, theme rels), shapes positioned outside the slide "
    "bounds, or a broken chart/picture reference. Do NOT re-issue identical code. "
    "Rebuild the content and prefer building fresh decks over round-tripping "
    "desktop-authored files."
)

_PPTX_SUFFIXES = (".pptx", ".pptm")

# 16:9 slide geometry (the canonical root size for every new deck).
_SLIDE_WIDTH_IN = 13.333
_SLIDE_HEIGHT_IN = 7.5
_SLIDE_WIDTH_EMU = 12192000
_SLIDE_HEIGHT_EMU = 6858000

_MAX_VIEW_SLIDES = 4
_RENDER_TIMEOUT_S = 120
_RENDER_DPI = 110


def init_tools_library(config: dict = None) -> None:
    try:
        import pipmaster as pm
        pm.ensure_packages(["python-pptx"])
    except Exception as e:
        ASCIIColors.warning(f"[powerpoint_tools] Failed to ensure dependencies: {e}")
    try:
        import pipmaster as pm
        pm.ensure_packages(["pypdfium2"])
    except Exception:
        pass  # best-effort: text fallback when unavailable


def _safe_import(name: Any, *args: Any, **kwargs: Any) -> Any:
    """Whitelist gate for recipe imports (installed as __import__ in the sandbox)."""
    module_name = str(name)
    root_mod = module_name.split(".")[0]
    if root_mod in _SAFE_STDLIB_MODULES or any(
        module_name == prefix.rstrip(".") or module_name.startswith(prefix)
        for prefix in _IMPORT_PREFIX_WHITELIST
    ):
        import importlib
        return importlib.import_module(module_name)
    raise ImportError(
        f"Import '{module_name}' is not allowed in recipes. Allowed imports: pptx.* "
        "and safe stdlib modules (math, datetime, random, json, statistics, itertools, "
        "functools, re, uuid, calendar, decimal, fractions, textwrap, base64, hashlib, "
        "string, collections). All design symbols are also pre-imported."
    )


def _build_safe_builtins() -> Dict[str, Any]:
    """Curated builtins for the recipe sandbox: no file/system/eval access."""
    import builtins
    base = {k: v for k, v in vars(builtins).items() if not k.startswith("_")}
    for dangerous in ("open", "eval", "exec", "compile", "breakpoint", "input",
                      "exit", "quit", "help", "license", "globals", "locals",
                      "vars", "getattr", "setattr", "delattr"):
        base.pop(dangerous, None)
    base["__import__"] = _safe_import
    return base


def _check_recipe(recipe: str) -> None:
    """Rejects recipes containing forbidden tokens before execution."""
    for token in _FORBIDDEN_TOKENS:
        if token in recipe:
            raise ValueError(
                f"Forbidden token '{token}' in recipe. Recipes are plain Python: imports "
                "are allowed only from the whitelist (pptx.* and safe stdlib modules), "
                "with no file/system access and no dunder access. Prefer the pre-imported "
                "design symbols."
            )


def _annotate_traceback(raw_traceback: str, source_code: str, source_name: str) -> str:
    """Appends the offending source line after each frame referencing source_name."""
    lines = source_code.splitlines()
    annotated_lines: List[str] = []
    for line in raw_traceback.splitlines():
        annotated_lines.append(line)
        match = re.search(re.escape(f'File "{source_name}", line ') + r"(\d+)", line)
        if match:
            line_no = int(match.group(1))
            if 1 <= line_no <= len(lines):
                annotated_lines.append(f"    >>> LINE {line_no}: {lines[line_no - 1].strip()}")
    return "\n".join(annotated_lines)


def _candidate_working_dirs(working_dir: Optional[Path]) -> List[Path]:
    """Builds the ordered list of directories scanned for media files."""
    candidates: List[Path] = []
    if working_dir is not None:
        try:
            candidates.append(Path(working_dir).resolve())
        except Exception:
            pass
    try:
        candidates.append(Path.cwd().resolve())
    except Exception:
        pass
    return candidates


def _resolve_media_path(ref: str, working_dir: Optional[Path]) -> Optional[Path]:
    """Resolves a media reference (name, relative or absolute path) to an existing file."""
    raw = (ref or "").strip()
    if not raw:
        return None
    cleaned = raw.replace("\\", "/")
    for prefix in ("./workspace/", "./data_workspace/", "workspace/", "data_workspace/"):
        if cleaned.lower().startswith(prefix):
            cleaned = cleaned[len(prefix):]
            break
    try:
        raw_path = Path(cleaned)
        if raw_path.is_absolute() and raw_path.is_file():
            return raw_path
    except Exception:
        pass
    for base in _candidate_working_dirs(working_dir):
        try:
            candidate = (base / cleaned).resolve()
            if candidate.is_file():
                return candidate
        except Exception:
            continue
    return None


def _list_media_files(working_dir: Optional[Path]) -> List[str]:
    """Lists media file names available to recipes (workspace-first, cwd fallback)."""
    seen: set = set()
    names: List[str] = []
    scanned = 0
    for base in _candidate_working_dirs(working_dir):
        if not base.exists() or not base.is_dir():
            continue
        try:
            for f in base.rglob("*"):
                scanned += 1
                if scanned > _MAX_MEDIA_SCAN:
                    break
                if not f.is_file() or f.suffix.lower() not in _MEDIA_EXTENSIONS:
                    continue
                rel = f.relative_to(base).as_posix()
                if rel not in seen:
                    seen.add(rel)
                    names.append(rel)
                if len(names) >= _MAX_MEDIA_LIST:
                    return names
        except Exception:
            continue
    return names


def _build_recipe_namespace(working_dir: Optional[Path] = None) -> Dict[str, Any]:
    """Pre-imports the python-pptx design vocabulary into the recipe namespace."""
    namespace: Dict[str, Any] = {}
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt, Emu
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.oxml.ns import qn

        namespace.update({
            "Presentation": Presentation,
            "Inches": Inches,
            "Pt": Pt,
            "Emu": Emu,
            "RGBColor": RGBColor,
            "MSO_SHAPE": MSO_SHAPE,
            "MSO_CONNECTOR": MSO_CONNECTOR,
            "MSO_ANCHOR": MSO_ANCHOR,
            "PP_ALIGN": PP_ALIGN,
            "qn": qn,
        })
    except ImportError as import_error:
        ASCIIColors.warning(
            f"[powerpoint_tools] python-pptx unavailable: {import_error}. "
            "Recipes will fail until 'pip install python-pptx' is run."
        )

    namespace["media"] = lambda ref: _resolve_media_path(ref, working_dir)
    namespace["list_media"] = lambda: _list_media_files(working_dir)
    return namespace


def _reopen_office_file(path: Path, suffix: str) -> Tuple[bool, str]:
    """Layer 3: re-opens the saved file with python-pptx (must contain >= 1 slide)."""
    try:
        from pptx import Presentation as _PptxPresentation
    except ImportError:
        return True, "ok (layers 1-2 only: python-pptx not installed)"
    try:
        prs = _PptxPresentation(str(path))
        if len(prs.slides) == 0:
            return False, "re-open succeeded but the presentation has zero slides"
        return True, "ok"
    except Exception as reopen_error:
        return False, f"python-pptx re-open failed: {reopen_error}"


def _validate_office_file(path: Path) -> Tuple[bool, str]:
    """Three-layer OOXML integrity validation for .pptx/.pptm files."""
    suffix = path.suffix.lower()
    validator = OOXML_VALIDATORS.get(suffix)
    if validator is None:
        return True, "ok (no validator for this suffix)"

    # Layer 1: ZIP container + required parts
    try:
        with zipfile.ZipFile(path, "r") as zf:
            names = set(zf.namelist())
            for required in validator["required_parts"]:
                if required not in names:
                    return False, f"missing required OOXML part: {required}"
    except zipfile.BadZipFile:
        return False, "file is not a valid ZIP/OOXML container"
    except Exception as zip_error:
        return False, f"ZIP inspection failed: {zip_error}"

    # Layer 2: strict XML parse of every .xml/.rels part
    try:
        with zipfile.ZipFile(path, "r") as zf:
            for name in zf.namelist():
                if not (name.endswith(".xml") or name.endswith(".rels")):
                    continue
                ET.fromstring(zf.read(name))
    except ET.ParseError as parse_error:
        return False, f"malformed XML part: {parse_error}"
    except Exception as xml_error:
        return False, f"XML inspection failed: {xml_error}"

    # Layer 3: library re-open
    return _reopen_office_file(path, suffix)


def _save_and_verify(obj: Any, save_path: Path) -> Dict[str, Any]:
    """Saves the object and verifies the result; rolls back to the pre-save bytes on failure."""
    try:
        snapshot = save_path.read_bytes() if save_path.exists() else None
    except Exception:
        snapshot = None

    try:
        obj.save(str(save_path))
    except Exception as save_error:
        if snapshot is not None:
            try:
                save_path.write_bytes(snapshot)
            except Exception:
                pass
        return {
            "success": False,
            "error": f"save() raised: {save_error}",
            "error_code": OFFICE_GUARD_FAILURE_CODE,
            "guidance": OFFICE_GUARD_GUIDANCE,
        }

    ok, detail = _validate_office_file(save_path)
    if ok:
        return {"success": True, "file_name": save_path.name}

    if snapshot is not None:
        try:
            save_path.write_bytes(snapshot)
        except Exception:
            pass
    ASCIIColors.error(f"[PowerPoint Guard] Validation failed for '{save_path.name}': {detail}")
    return {
        "success": False,
        "error": f"saved file failed integrity validation: {detail}",
        "error_code": OFFICE_GUARD_FAILURE_CODE,
        "guidance": OFFICE_GUARD_GUIDANCE,
    }


def _exec_recipe(recipe: str, namespace: Dict[str, Any], source_name: str) -> Optional[Dict[str, Any]]:
    """Executes a recipe inside the sandbox namespace; returns an error dict on failure."""
    if len(recipe) > _MAX_RECIPE_CHARS:
        return {
            "success": False,
            "error": f"recipe exceeds {_MAX_RECIPE_CHARS} characters ({len(recipe)} given).",
        }
    try:
        _check_recipe(recipe)
    except ValueError as check_error:
        return {"success": False, "error": str(check_error)}

    namespace["__builtins__"] = _build_safe_builtins()
    try:
        exec(compile(textwrap.dedent(recipe), source_name, "exec"), namespace)
        return None
    except Exception:
        annotated = _annotate_traceback(traceback.format_exc(), textwrap.dedent(recipe), source_name)
        if "save(" in annotated and "missing" in annotated:
            annotated += (
                "\nHint: do NOT call prs.save() — the harness saves and verifies "
                "the file automatically after your recipe runs."
            )
        return {"success": False, "error": annotated}


def _resolve_working_dir(discussion_instance: Any) -> Optional[Path]:
    """Resolves the discussion workspace root for media resolution."""
    if discussion_instance is None:
        return None
    ws_data = getattr(discussion_instance, "workspace_data_path", None)
    if ws_data:
        try:
            return Path(ws_data).resolve()
        except Exception:
            return None
    base_ws = getattr(discussion_instance, "workspace_path", None)
    disc_id = getattr(discussion_instance, "id", None)
    if base_ws and disc_id:
        try:
            return (Path(base_ws) / str(disc_id) / "workspace_data").resolve()
        except Exception:
            return None
    return None


def _sanitize_file_name(file: str, default: str) -> str:
    """Normalizes an LLM-provided file name into a safe relative workspace path."""
    cleaned = (file or "").replace("\\", "/").strip().strip("/")
    cleaned = re.sub(r"^[A-Za-z]:", "", cleaned)
    parts: List[str] = []
    for seg in cleaned.split("/"):
        seg = seg.strip()
        if not seg or seg in (".", ".."):
            continue
        seg = re.sub(r'[<>:"|?*\x00-\x1f]', "", seg).strip()
        if seg:
            parts.append(seg)
    if not parts:
        return default
    return "/".join(parts)


def _build_result(save_result: Dict[str, Any], save_path: Path) -> Dict[str, Any]:
    """Shapes the final tool result from a save/verify outcome."""
    if save_result.get("success"):
        return {
            "success": True,
            "file_name": save_path.name,
            "workspace_file": save_path.name,
            "output": f"PowerPoint file '{save_path.name}' saved and verified.",
        }
    return {
        "success": False,
        "error": save_result.get("error", "save failed"),
        "error_code": save_result.get("error_code"),
        "guidance": save_result.get("guidance"),
        "output": save_result.get("error", "save failed"),
    }


def _require_suffix(file: str, allowed_suffixes: Tuple[str, ...], tool_name: str) -> Optional[Dict[str, Any]]:
    """Rejects files whose extension is not in the allowed family."""
    suffix = Path(file).suffix.lower()
    if suffix not in allowed_suffixes:
        return {
            "success": False,
            "error": (
                f"'{file}' is not a supported file for {tool_name}. "
                f"Allowed extensions: {', '.join(allowed_suffixes)}."
            ),
        }
    return None


def _create_root_object(suffix: str, namespace: Dict[str, Any]) -> Any:
    """Creates a fresh 16:9 presentation and binds it as the recipe root."""
    prs = namespace["Presentation"]()
    prs.slide_width = namespace["Emu"](_SLIDE_WIDTH_EMU)
    prs.slide_height = namespace["Emu"](_SLIDE_HEIGHT_EMU)
    namespace["prs"] = prs
    namespace["presentation"] = prs
    return prs


def _open_existing(path: Path, namespace: Dict[str, Any]) -> Any:
    """Opens an existing deck and binds it as the recipe root."""
    prs = namespace["Presentation"](str(path))
    namespace["prs"] = prs
    namespace["presentation"] = prs
    return prs


def _pick_effective_root(namespace: Dict[str, Any], suffix: str, original_root: Any) -> Any:
    """Detects a recipe-created root bound to a known alias (root-rebind detection)."""
    for alias in ("prs", "presentation"):
        candidate = namespace.get(alias)
        if candidate is not None and candidate is not original_root:
            return candidate
    return original_root


def _read_pptx(path: Path) -> str:
    """Builds a markdown inventory of every slide (title + shape texts)."""
    try:
        from pptx import Presentation as _PptxPresentation
    except ImportError:
        return "python-pptx is not installed; cannot extract the deck inventory."

    prs = _PptxPresentation(str(path))
    lines: List[str] = [f"# Deck inventory: {path.name}", ""]
    if len(prs.slides) == 0:
        lines.append("(the deck contains no slides)")
        return "\n".join(lines)

    for idx, slide in enumerate(prs.slides, start=1):
        lines.append(f"## Slide {idx}")
        title = None
        try:
            title = slide.shapes.title
        except Exception:
            title = None
        if title is not None and (title.text or "").strip():
            lines.append(f"- Title: {title.text.strip()}")
        for shape in slide.shapes:
            if shape is title:
                continue
            text = ""
            try:
                text = (shape.text or "").strip()
            except Exception:
                text = ""
            if text:
                first_line = text.splitlines()[0]
                lines.append(f"- {shape.shape_type}: {first_line[:160]}")
        lines.append("")
    return "\n".join(lines)


def _parse_slides_spec(spec: str, deck_size: int) -> List[int]:
    """Parses a slides spec ('1-4', '2,5,7', '3') into a deduplicated, clamped, capped list."""
    raw = (spec or "").strip()
    if not raw:
        raw = "1-4"
    picked: List[int] = []
    for chunk in raw.split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        range_match = re.match(r"^(\d+)\s*-\s*(\d+)$", chunk)
        if range_match:
            start = max(1, int(range_match.group(1)))
            end = min(deck_size, int(range_match.group(2)))
            for n in range(start, end + 1):
                picked.append(n)
            continue
        if chunk.isdigit():
            n = int(chunk)
            if 1 <= n <= deck_size:
                picked.append(n)
            continue
    deduped: List[int] = []
    seen: set = set()
    for n in picked:
        if n not in seen:
            seen.add(n)
            deduped.append(n)
    return deduped[:_MAX_VIEW_SLIDES]


def _render_slides_to_pngs(deck_path: Path, slide_numbers: List[int]) -> Optional[List[bytes]]:
    """Renders the requested slides to PNG bytes via soffice + pypdfium2; None on unavailability."""
    soffice = shutil.which("soffice")
    if soffice is None:
        return None
    try:
        import pypdfium2 as pdfium  # noqa: F401
    except ImportError:
        return None

    tmp_dir = Path(tempfile.mkdtemp(prefix="pptx_view_"))
    try:
        cmd = [
            soffice, "--headless", "--convert-to", "pdf",
            "--outdir", str(tmp_dir), str(deck_path),
        ]
        completed = subprocess.run(
            cmd, capture_output=True, timeout=_RENDER_TIMEOUT_S,
        )
        pdf_path = tmp_dir / (deck_path.stem + ".pdf")
        if completed.returncode != 0 or not pdf_path.exists():
            return None

        import pypdfium2 as pdfium
        png_bytes: List[bytes] = []
        pdf = pdfium.PdfDocument(str(pdf_path))
        try:
            for slide_no in slide_numbers:
                page_idx = slide_no - 1
                if page_idx >= len(pdf):
                    continue
                page = pdf[page_idx]
                bitmap = page.render(scale=_RENDER_DPI / 72.0)
                pil_image = bitmap.to_pil()
                import io as _io
                buf = _io.BytesIO()
                pil_image.save(buf, format="PNG")
                png_bytes.append(buf.getvalue())
        finally:
            pdf.close()
        return png_bytes if png_bytes else None
    except Exception:
        return None
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _slides_text_fallback(deck_path: Path, slide_numbers: List[int]) -> str:
    """Structural slide-by-slide inventory used when visual rendering is unavailable."""
    try:
        from pptx import Presentation as _PptxPresentation
    except ImportError:
        return "python-pptx is not installed; cannot extract slide contents."

    prs = _PptxPresentation(str(deck_path))
    lines: List[str] = []
    for slide_no in slide_numbers:
        if slide_no > len(prs.slides):
            continue
        slide = prs.slides[slide_no - 1]
        lines.append(f"## Slide {slide_no}")
        for shape in slide.shapes:
            try:
                text = (shape.text or "").strip()
            except Exception:
                text = ""
            if text:
                lines.append(f"- {text.splitlines()[0][:200]}")
        lines.append("")
    if not lines:
        return "No content found for the requested slides."
    return "\n".join(lines)


def tool_build_powerpoint(
    file: str = "presentation.pptx",
    recipe: str = "",
    discussion_instance=None,
    lollms_client_instance=None,
) -> Dict[str, Any]:
    """Builds a new PowerPoint deck from sandboxed python-pptx recipe code.

    The root presentation `prs` is pre-bound and set to 16:9 (13.333in x 7.5in).
    Discover the available slide layouts with tool_pptx_info BEFORE building and
    add slides via `prs.slide_layouts[i]`. Imports are allowed ONLY from the
    whitelist (pptx.* plus safe stdlib modules). Do NOT call prs.save() — the
    file is saved and integrity-verified automatically. Verify visually with
    tool_view_slides (max 4 slides per call, transient images).

    Args:
        file (str): Output file name (default 'presentation.pptx'; .pptx or .pptm).
        recipe (str): Python recipe code manipulating the pre-bound `prs` root.

    Returns:
        dict: {"success", "file_name", "workspace_file", "output"} on success;
        {"success": False, "error", "guidance"} on failure.
    """
    suffix_check = _require_suffix(file, _PPTX_SUFFIXES, "tool_build_powerpoint")
    if suffix_check:
        return suffix_check

    file_name = _sanitize_file_name(file, "presentation.pptx")
    working_dir = _resolve_working_dir(discussion_instance)
    base_dir = working_dir if working_dir is not None else Path.cwd()
    save_path = base_dir / file_name

    namespace = _build_recipe_namespace(working_dir)
    if "Presentation" not in namespace:
        return {
            "success": False,
            "error": "python-pptx is not installed. Run 'pip install python-pptx' first.",
        }

    try:
        root_obj = _create_root_object(Path(file_name).suffix.lower(), namespace)
    except Exception as root_error:
        return {"success": False, "error": f"failed to create root presentation: {root_error}"}

    # No-op save shim: the harness owns persistence; recipes must not save.
    try:
        root_obj.save = lambda *a, **k: None  # type: ignore[method-assign]
    except Exception:
        pass

    exec_error = _exec_recipe(recipe, namespace, "<pptx_recipe>")
    if exec_error is not None:
        return exec_error

    effective_root = _pick_effective_root(namespace, Path(file_name).suffix.lower(), root_obj)
    save_result = _save_and_verify(effective_root, save_path)
    return _build_result(save_result, save_path)


def tool_edit_powerpoint(
    file: str,
    recipe: str = "",
    discussion_instance=None,
    lollms_client_instance=None,
) -> Dict[str, Any]:
    """Edits an existing PowerPoint deck with sandboxed python-pptx recipe code.

    The existing deck is opened and pre-bound as `prs`. Imports are allowed ONLY
    from the whitelist (pptx.* plus safe stdlib modules). Do NOT call prs.save() —
    the file is saved and integrity-verified automatically; the pre-edit bytes are
    restored if validation fails. Verify visually with tool_view_slides (max 4
    slides per call, transient images).

    Args:
        file (str): Deck file name to edit (.pptx or .pptm).
        recipe (str): Python recipe code manipulating the pre-bound `prs` root.

    Returns:
        dict: {"success", "file_name", "workspace_file", "output"} on success;
        {"success": False, "error", "guidance"} on failure.
    """
    suffix_check = _require_suffix(file, _PPTX_SUFFIXES, "tool_edit_powerpoint")
    if suffix_check:
        return suffix_check

    file_name = _sanitize_file_name(file, "presentation.pptx")
    working_dir = _resolve_working_dir(discussion_instance)
    base_dir = working_dir if working_dir is not None else Path.cwd()
    save_path = base_dir / file_name
    if not save_path.exists():
        return {"success": False, "error": f"file '{file_name}' not found in the workspace."}

    namespace = _build_recipe_namespace(working_dir)
    if "Presentation" not in namespace:
        return {
            "success": False,
            "error": "python-pptx is not installed. Run 'pip install python-pptx' first.",
        }

    try:
        root_obj = _open_existing(save_path, namespace)
    except Exception as open_error:
        return {"success": False, "error": f"failed to open '{file_name}': {open_error}"}

    try:
        root_obj.save = lambda *a, **k: None  # type: ignore[method-assign]
    except Exception:
        pass

    exec_error = _exec_recipe(recipe, namespace, "<pptx_recipe>")
    if exec_error is not None:
        return exec_error

    effective_root = _pick_effective_root(namespace, Path(file_name).suffix.lower(), root_obj)
    save_result = _save_and_verify(effective_root, save_path)
    return _build_result(save_result, save_path)


def tool_read_powerpoint(
    file: str,
    discussion_instance=None,
    lollms_client_instance=None,
) -> Dict[str, Any]:
    """Reads a PowerPoint deck and returns a markdown inventory of its slides.

    Validates the file first (three-layer OOXML check), then extracts each
    slide's title and shape texts. For layout/template details use
    tool_pptx_info; for rendered pixels use tool_view_slides.

    Args:
        file (str): Deck file name (.pptx or .pptm).

    Returns:
        dict: {"success", "file_name", "workspace_file", "inventory"}.
    """
    suffix_check = _require_suffix(file, _PPTX_SUFFIXES, "tool_read_powerpoint")
    if suffix_check:
        return suffix_check

    file_name = _sanitize_file_name(file, "presentation.pptx")
    working_dir = _resolve_working_dir(discussion_instance)
    base_dir = working_dir if working_dir is not None else Path.cwd()
    save_path = base_dir / file_name
    if not save_path.exists():
        return {"success": False, "error": f"file '{file_name}' not found in the workspace."}

    ok, detail = _validate_office_file(save_path)
    if not ok:
        return {"success": False, "error": f"file failed integrity validation: {detail}"}

    inventory = _read_pptx(save_path)
    return {
        "success": True,
        "file_name": save_path.name,
        "workspace_file": save_path.name,
        "inventory": inventory,
        "output": inventory,
    }


def tool_pptx_info(
    file: str = "presentation.pptx",
    discussion_instance=None,
    lollms_client_instance=None,
) -> Dict[str, Any]:
    """Returns template/size awareness for a deck: slide size, layout and placeholder table.

    MANDATORY AWARENESS STEP: call this BEFORE building to discover the deck's
    slide size (inches + EMU + aspect ratio), every slide layout name with its
    placeholders (idx + type), and the theme fonts (best-effort). New decks are
    created at 16:9 (13.333in x 7.5in).

    Args:
        file (str): Deck file name (.pptx or .pptm). A non-existent file reports
            the canonical 16:9 root geometry for new decks.

    Returns:
        dict: {"success", "output" (markdown), "file_name", "workspace_file"}.
    """
    file_name = _sanitize_file_name(file, "presentation.pptx")
    working_dir = _resolve_working_dir(discussion_instance)
    base_dir = working_dir if working_dir is not None else Path.cwd()
    deck_path = base_dir / file_name

    lines: List[str] = [f"# PowerPoint template info: {file_name}", ""]

    if not deck_path.exists():
        lines.append(
            "The file does not exist yet. New decks are created at 16:9 — "
            f"{_SLIDE_WIDTH_IN}in x {_SLIDE_HEIGHT_IN}in "
            f"(EMU {_SLIDE_WIDTH_EMU} x {_SLIDE_HEIGHT_EMU})."
        )
        return {
            "success": True,
            "output": "\n".join(lines),
            "file_name": file_name,
            "workspace_file": file_name,
        }

    try:
        from pptx import Presentation as _PptxPresentation
        from pptx.util import Emu as _Emu
    except ImportError:
        return {
            "success": False,
            "error": "python-pptx is not installed. Run 'pip install python-pptx' first.",
        }

    try:
        prs = _PptxPresentation(str(deck_path))
    except Exception as open_error:
        return {"success": False, "error": f"failed to open '{file_name}': {open_error}"}

    width_emu = int(prs.slide_width)
    height_emu = int(prs.slide_height)
    width_in = width_emu / 914400.0
    height_in = height_emu / 914400.0
    aspect = round(width_in / height_in, 3) if height_in else None

    lines.append("## Slide size")
    lines.append(f"- {width_in:.3f}in x {height_in:.3f}in (EMU {width_emu} x {height_emu})")
    if aspect is not None:
        lines.append(f"- Aspect ratio: {aspect}")
    lines.append(f"- Slide count: {len(prs.slides)}")
    lines.append("")

    lines.append("## Slide layouts (prs.slide_layouts[i])")
    for i, layout in enumerate(prs.slide_layouts):
        lines.append(f"### [{i}] {layout.name}")
        placeholders = list(layout.placeholders)
        if not placeholders:
            lines.append("- (no placeholders)")
            continue
        for ph in placeholders:
            ph_idx = getattr(ph, "placeholder_format").idx
            ph_type = str(getattr(ph, "placeholder_format").type)
            lines.append(f"- idx={ph_idx}, type={ph_type}")
        lines.append("")

    # Theme fonts (best-effort).
    try:
        theme_fonts: List[str] = []
        for layout in prs.slide_layouts:
            master = getattr(layout, "slide_master", None)
            if master is None:
                continue
            theme = getattr(master, "element", None)
            if theme is None:
                continue
            for el in theme.iter():
                tag = el.tag.lower()
                if tag.endswith("}majorfont") or tag.endswith("}minorfont"):
                    for child in el:
                        child_tag = child.tag.lower()
                        if child_tag.endswith("}latin"):
                            face = child.get("typeface")
                            if face:
                                label = "major" if "majorfont" in tag else "minor"
                                entry = f"{label}: {face}"
                                if entry not in theme_fonts:
                                    theme_fonts.append(entry)
        if theme_fonts:
            lines.append("## Theme fonts (best-effort)")
            for entry in theme_fonts:
                lines.append(f"- {entry}")
    except Exception:
        pass

    return {
        "success": True,
        "output": "\n".join(lines),
        "file_name": file_name,
        "workspace_file": file_name,
    }


def tool_view_slides(
    file: str,
    slides: str = "1-4",
    discussion_instance=None,
    lollms_client_instance=None,
) -> Dict[str, Any]:
    """Renders up to 4 slides of a deck to images attached to THIS ROUND ONLY.

    TRANSIENT VISION: the rendered slide images are attached to the current
    generation round as vision content and are NEVER stored in the workspace.
    A maximum of 4 slides per call is enforced (extra slides are truncated with
    a note). Rendering requires LibreOffice (soffice) and pypdfium2; when either
    is unavailable, a structural text inventory of the requested slides is
    returned instead.

    Args:
        file (str): Deck file name (.pptx or .pptm).
        slides (str): Slide selection — '1-4' (range), '2,5,7' (list) or '3' (single).

    Returns:
        dict: {"success", "output", "images" (b64 PNG list), "image_media_types",
        "prompt_injection"} — images flow to the current round only.
    """
    suffix_check = _require_suffix(file, _PPTX_SUFFIXES, "tool_view_slides")
    if suffix_check:
        return suffix_check

    file_name = _sanitize_file_name(file, "presentation.pptx")
    working_dir = _resolve_working_dir(discussion_instance)
    base_dir = working_dir if working_dir is not None else Path.cwd()
    deck_path = base_dir / file_name
    if not deck_path.exists():
        return {"success": False, "error": f"file '{file_name}' not found in the workspace."}

    try:
        from pptx import Presentation as _PptxPresentation
        prs = _PptxPresentation(str(deck_path))
        deck_size = len(prs.slides)
    except Exception as open_error:
        return {"success": False, "error": f"failed to open '{file_name}': {open_error}"}

    if deck_size == 0:
        return {"success": False, "error": "the deck contains no slides to render."}

    requested_raw = (slides or "").strip() or "1-4"
    slide_numbers = _parse_slides_spec(requested_raw, deck_size)
    if not slide_numbers:
        return {
            "success": False,
            "error": f"no valid slides in spec '{requested_raw}' (deck has {deck_size} slides).",
        }

    truncated_note = ""
    all_requested = [
        n for n in _parse_slides_spec_unlimited(requested_raw, deck_size)
    ] if False else []
    # Compute the full requested list only to report truncation.
    full_requested: List[int] = []
    seen: set = set()
    for chunk in requested_raw.split(","):
        chunk = chunk.strip()
        range_match = re.match(r"^(\d+)\s*-\s*(\d+)$", chunk)
        if range_match:
            start = max(1, int(range_match.group(1)))
            end = min(deck_size, int(range_match.group(2)))
            for n in range(start, end + 1):
                if n not in seen:
                    seen.add(n)
                    full_requested.append(n)
        elif chunk.isdigit():
            n = int(chunk)
            if 1 <= n <= deck_size and n not in seen:
                seen.add(n)
                full_requested.append(n)
    if len(full_requested) > _MAX_VIEW_SLIDES:
        truncated_note = (
            f" (requested {len(full_requested)} slides; only the first "
            f"{_MAX_VIEW_SLIDES} are attached per call)"
        )

    png_bytes = _render_slides_to_pngs(deck_path, slide_numbers)
    if png_bytes:
        b64_images = [base64.b64encode(png).decode("utf-8") for png in png_bytes]
        slide_list = ", ".join(str(n) for n in slide_numbers)
        output_text = (
            f"Attached {len(b64_images)} slide image(s) (slides {slide_list}) "
            f"rendered from '{file_name}'{truncated_note}."
        )
        return {
            "success": True,
            "output": output_text,
            "images": b64_images,
            "image_media_types": ["image/png"] * len(b64_images),
            "prompt_injection": (
                f"\n\n[SYSTEM: {len(b64_images)} slide image(s) attached as vision content "
                "for THIS ROUND ONLY. Use them to verify layout/content. They are NOT "
                "stored in the workspace.]"
            ),
        }

    fallback_text = _slides_text_fallback(deck_path, slide_numbers)
    output_text = (
        f"Visual rendering unavailable (LibreOffice/pypdfium2 missing) — returning "
        f"structural inventory instead for slides "
        f"{', '.join(str(n) for n in slide_numbers)} of '{file_name}'{truncated_note}:\n\n"
        f"{fallback_text}"
    )
    return {
        "success": True,
        "output": output_text,
        "prompt_injection": (
            "\n\n[SYSTEM: Visual slide rendering is unavailable in this environment. "
            "A structural text inventory was returned instead.]"
        ),
    }


def get_tools() -> List[Dict[str, Any]]:
    """Returns the five PowerPoint tool specifications for LCP registration."""
    return [
        {
            "name": "tool_build_powerpoint",
            "description": (
                "Build a new PowerPoint deck (.pptx/.pptm) from sandboxed python-pptx "
                "recipe code. The root `prs` is pre-bound at 16:9 (13.333in x 7.5in). "
                "Discover layouts with tool_pptx_info BEFORE building. Imports only from "
                "pptx.* + safe stdlib. Do NOT call prs.save() — auto-saved and verified."
            ),
            "parameters": [
                {"name": "file", "type": "string", "description": "Output file name (.pptx or .pptm)."},
                {"name": "recipe", "type": "string", "description": "Python recipe manipulating the pre-bound `prs` root."},
            ],
            "output": "JSON result with success, file_name, workspace_file.",
            "callable": tool_build_powerpoint,
            "_source": "powerpoint_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_edit_powerpoint",
            "description": (
                "Edit an existing PowerPoint deck with sandboxed python-pptx recipe code. "
                "The deck is pre-bound as `prs`. Imports only from pptx.* + safe stdlib. "
                "Do NOT call prs.save() — auto-saved, verified, rolled back on failure."
            ),
            "parameters": [
                {"name": "file", "type": "string", "description": "Deck file name to edit (.pptx or .pptm)."},
                {"name": "recipe", "type": "string", "description": "Python recipe manipulating the pre-bound `prs` root."},
            ],
            "output": "JSON result with success, file_name, workspace_file.",
            "callable": tool_edit_powerpoint,
            "_source": "powerpoint_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_read_powerpoint",
            "description": (
                "Read a PowerPoint deck and return a markdown inventory of its slides "
                "(titles + shape texts). Validated with a three-layer OOXML check first."
            ),
            "parameters": [
                {"name": "file", "type": "string", "description": "Deck file name (.pptx or .pptm)."},
            ],
            "output": "JSON result with success, file_name, inventory.",
            "callable": tool_read_powerpoint,
            "_source": "powerpoint_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_pptx_info",
            "description": (
                "Template/size awareness: slide size (inches + EMU + aspect), slide count, "
                "per-layout placeholder table (idx + type) and theme fonts. Call BEFORE "
                "building to pick the right layouts via prs.slide_layouts[i]."
            ),
            "parameters": [
                {"name": "file", "type": "string", "description": "Deck file name (.pptx or .pptm)."},
            ],
            "output": "JSON result with success, output (markdown).",
            "callable": tool_pptx_info,
            "_source": "powerpoint_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_view_slides",
            "description": (
                "Visually verify a deck: renders up to 4 slides per call as images attached "
                "to the CURRENT ROUND ONLY (never stored in the workspace). Requires "
                "LibreOffice + pypdfium2; falls back to a structural text inventory."
            ),
            "parameters": [
                {"name": "file", "type": "string", "description": "Deck file name (.pptx or .pptm)."},
                {"name": "slides", "type": "string", "description": "Slide selection: '1-4', '2,5,7' or '3'."},
            ],
            "output": "JSON result with success, output, images (transient).",
            "callable": tool_view_slides,
            "_source": "powerpoint_tools",
            "_source_file": __file__,
        },
    ]