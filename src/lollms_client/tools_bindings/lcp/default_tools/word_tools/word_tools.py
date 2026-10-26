"""
Word Tools — dedicated LCP toolset for Microsoft Word documents.

Doctrine (read this before writing any recipe):
- Full edit freedom: the LLM may use its own coding skills with python-docx.
  These are HINTS on what is possible, not restrictions.
- Comments: use the pre-bound add_comment(doc, paragraph, text, author, run)
  helper (real OXML comments injected into comments.xml — the comments part is
  created automatically when missing). Pass run=resolve_comment_runs(doc,
  "text") to anchor the comment on the exact text fragment.
- Highlights: run.font.highlight_color = WD_COLOR_INDEX.YELLOW.
- Reformat: styles, fonts, alignment (WD_ALIGN_PARAGRAPH), spacing
  (WD_LINE_SPACING), colors (DocxRGBColor).
- Sections: doc.add_section(WD_SECTION.NEW_PAGE).
- Links/fields: python-docx cannot refresh external links; insert field codes
  (TOC, HYPERLINK) via OXML using the pre-bound qn() helper.
- Review: list_comments(doc) / list_highlights(doc) expose existing review
  state to recipes; tool_word_review reads it without a recipe.
"""
import re
import textwrap
import traceback
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from ascii_colors import ASCIIColors

TOOL_LIBRARY_NAME = "Word Tools"
TOOL_LIBRARY_DESC = (
    "Create, edit and read Word documents (.docx/.docm) via sandboxed python-docx "
    "recipe code. Full edit vocabulary: REAL anchored comments (add_comment helper "
    "auto-creates the comments part and anchors commentRangeStart/End on the exact "
    "run via resolve_comment_runs), highlights (WD_COLOR_INDEX), reformatting "
    "(styles/fonts/alignment/spacing), sections (WD_SECTION), links/fields via OXML "
    "(qn), plus review helpers (list_comments, list_highlights) and a dedicated "
    "tool_word_review reader. Hints, not restrictions — the LLM codes freely. Every "
    "save is verified with a strict three-layer OOXML integrity check before being "
    "committed to the workspace."
)
TOOL_LIBRARY_ICON = "📝"

_MAX_RECIPE_CHARS = 60_000
_FORBIDDEN_TOKENS = (
    "eval(", "exec(", "compile(", "open(",
    "os.", "sys.", "subprocess", "shutil", "pathlib", "__",
    "globals(", "locals(", "vars(", "getattr(", "setattr(", "delattr(",
)

_IMPORT_PREFIX_WHITELIST = ("docx.",)
_SAFE_STDLIB_MODULES = frozenset({
    "math", "datetime", "random", "json", "statistics", "itertools",
    "functools", "re", "uuid", "calendar", "decimal", "fractions",
    "textwrap", "base64", "hashlib", "string", "collections",
})

_MEDIA_EXTENSIONS = (
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg", ".tiff", ".ico",
    ".mp4", ".mov", ".avi", ".mkv", ".webm", ".wmv",
)
_MAX_MEDIA_SCAN = 5000
_MAX_MEDIA_LIST = 2000

OOXML_VALIDATORS = {
    ".docx": {"required_parts": ("[Content_Types].xml", "word/document.xml", "word/_rels/document.xml.rels"), "reopener": "python-docx"},
    ".docm": {"required_parts": ("[Content_Types].xml", "word/document.xml", "word/_rels/document.xml.rels"), "reopener": "python-docx"},
}
GUARDED_EXTENSIONS = tuple(OOXML_VALIDATORS.keys())

OFFICE_GUARD_FAILURE_CODE = "office_file_failed_validation"

OFFICE_GUARD_GUIDANCE = (
    "Likely causes: python-docx re-saves dropping parts desktop Word requires "
    "(styles.xml, numbering rels, theme rels), malformed OXML injected via qn(), "
    "or a broken comments part. Do NOT re-issue identical code. Rebuild the "
    "content: prefer building fresh documents over round-tripping desktop-authored "
    "files, and validate any OXML element you construct."
)

_DOCX_SUFFIXES = (".docx", ".docm")


def init_tools_library(config: dict = None) -> None:
    try:
        import pipmaster as pm
        pm.ensure_packages(["python-docx"])
    except Exception as e:
        ASCIIColors.warning(f"[word_tools] Failed to ensure dependencies: {e}")


def _safe_import(name: Any, *args: Any, **kwargs: Any) -> Any:
    """Whitelist gate for recipe imports (installed as __import__ in the sandbox)."""
    module_name = str(name)
    root_mod = module_name.split(".")[0]
    if root_mod in _SAFE_STDLIB_MODULES or any(
        module_name == prefix.rstrip(".") or module_name.startswith(prefix)
        for prefix in _IMPORT_PREFIX_WHITELIST
    ):
        import importlib
        return importlib.import_module(module_name)
    raise ImportError(
        f"Import '{module_name}' is not allowed in recipes. Allowed imports: docx.* "
        "and safe stdlib modules (math, datetime, random, json, statistics, itertools, "
        "functools, re, uuid, calendar, decimal, fractions, textwrap, base64, hashlib, "
        "string, collections). All design symbols are also pre-imported."
    )


def _build_safe_builtins() -> Dict[str, Any]:
    """Curated builtins for the recipe sandbox: no file/system/eval access."""
    import builtins
    base = {k: v for k, v in vars(builtins).items() if not k.startswith("_")}
    for dangerous in ("open", "eval", "exec", "compile", "breakpoint", "input",
                      "exit", "quit", "help", "license", "globals", "locals",
                      "vars", "getattr", "setattr", "delattr"):
        base.pop(dangerous, None)
    base["__import__"] = _safe_import
    return base


def _check_recipe(recipe: str) -> None:
    """Rejects recipes containing forbidden tokens before execution."""
    for token in _FORBIDDEN_TOKENS:
        if token in recipe:
            raise ValueError(
                f"Forbidden token '{token}' in recipe. Recipes are plain Python: imports "
                "are allowed only from the whitelist (docx.* and safe stdlib modules), "
                "with no file/system access and no dunder access. Prefer the pre-imported "
                "design symbols."
            )


def _annotate_traceback(raw_traceback: str, source_code: str, source_name: str) -> str:
    """Appends the exact failing recipe line after each frame referencing source_name."""
    code_lines = source_code.splitlines()
    annotated_lines = []
    for line in raw_traceback.splitlines():
        annotated_lines.append(line)
        match = re.search(
            re.escape(f'File "{source_name}", line ') + r"(\d+)", line
        )
        if match:
            failing_line_no = int(match.group(1))
            if 1 <= failing_line_no <= len(code_lines):
                annotated_lines.append(
                    f"    >>> LINE {failing_line_no}: {code_lines[failing_line_no - 1]}"
                )
    return "\n".join(annotated_lines)


def _candidate_working_dirs(working_dir: Optional[Path]) -> List[Path]:
    """Enumerates the roots a media reference may be anchored to.

    Workspace-first anchoring: when a discussion workspace is provided it is
    the SOLE root, so bare filenames resolve against the workspace and never
    against the process CWD (the daemon's startup folder). The CWD fallback
    applies only when no workspace was supplied (standalone/test harness).
    """
    if working_dir is not None:
        try:
            resolved = Path(working_dir).resolve()
        except Exception:
            resolved = None
        if resolved is not None and resolved.is_dir():
            return [resolved]
    try:
        cwd = Path(".").resolve()
    except Exception:
        return []
    return [cwd] if cwd.is_dir() else []


def _resolve_media_path(ref: str, working_dir: Optional[Path]) -> Path:
    """Resolves a workspace-relative media reference to an absolute path (read-only)."""
    raw = str(ref or "").strip().strip('"').strip("'")
    if not raw:
        raise ValueError("media() received an empty reference.")
    if "\\" in raw:
        raw = raw.replace("\\", "/")
    while raw.startswith("./"):
        raw = raw[2:]
    if Path(raw).is_absolute() or raw.startswith("/"):
        raise ValueError(f"Media reference '{ref}' must be workspace-relative, not absolute.")
    if ".." in Path(raw).parts:
        raise ValueError(f"Media reference '{ref}' must not contain '..' (workspace escape rejected).")

    roots = _candidate_working_dirs(working_dir)
    if not roots:
        raise ValueError(f"Media file '{ref}' could not be resolved: no accessible workspace directory.")

    target_name = Path(raw).name
    searched_roots: List[str] = []
    for root in roots:
        searched_roots.append(str(root))
        candidate = (root / raw).resolve()
        if candidate.is_relative_to(root) and candidate.is_file():
            return candidate
    for root in roots:
        scanned = 0
        for hit in root.rglob(target_name):
            scanned += 1
            if scanned > _MAX_MEDIA_SCAN:
                break
            hit_resolved = hit.resolve()
            if hit_resolved.is_relative_to(root) and hit_resolved.is_file():
                return hit_resolved
    available = []
    for root in roots:
        available.extend(_list_media_files(root))
        if len(available) >= 20:
            break
    available_hint = ", ".join(available[:20]) if available else "(none found)"
    raise ValueError(
        f"Media file '{ref}' not found. Searched roots: {', '.join(searched_roots)} "
        f"(exact path and recursive name search, subfolders included). "
        f"Available media: {available_hint}. Use list_media() for the full list."
    )


def _list_media_files(working_dir: Path) -> List[str]:
    """Enumerates media files under the workspace (subfolders included). Never raises."""
    root = Path(working_dir).resolve()
    if not root.is_dir():
        return []
    found: List[str] = []
    scanned = 0
    for entry in root.rglob("*"):
        scanned += 1
        if scanned > _MAX_MEDIA_SCAN:
            break
        if entry.suffix.lower() in _MEDIA_EXTENSIONS:
            try:
                rel = entry.resolve().relative_to(root)
            except ValueError:
                continue
            found.append(rel.as_posix())
            if len(found) >= _MAX_MEDIA_LIST:
                break
    return sorted(found, key=str.lower)


def _ensure_comments_part(doc: Any) -> Any:
    """Returns the document's comments part, creating and wiring it when absent.

    Mirrors the proven OPC part-creation pattern: a Part with an explicit
    comments content type at /word/comments.xml, related to the main document
    part via the COMMENTS relationship. The part is then included in package
    saves and [Content_Types].xml automatically.
    """
    for rel in doc.part.rels.values():
        if "comments" in rel.reltype:
            return rel.target_part

    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    comments_part = Part(
        PackURI("/word/comments.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        parse_xml(f'<w:comments {nsdecls("w")}/>'),
        doc.part.package,
    )
    doc.part.relate_to(comments_part, RT.COMMENTS)
    return comments_part


def _list_comments(doc: Any) -> List[Dict[str, Any]]:
    """Read-only extraction of all comments (author/initials/date/text). Never raises."""
    try:
        from docx.oxml.ns import qn

        comments: List[Dict[str, Any]] = []
        for rel in doc.part.rels.values():
            if "comments" not in rel.reltype:
                continue
            for node in rel.target_part.element.findall(qn("w:comment")):
                try:
                    text = "".join(
                        t.text or "" for t in node.findall(f".//{qn('w:t')}")
                    )
                    comments.append({
                        "id": node.get(qn("w:id")),
                        "author": node.get(qn("w:author")) or "",
                        "initials": node.get(qn("w:initials")) or "",
                        "date": node.get(qn("w:date")) or "",
                        "text": text,
                    })
                except Exception:
                    continue
            break
        return comments
    except Exception:
        return []


def _list_highlights(doc: Any, max_entries: int = 200) -> List[Dict[str, Any]]:
    """Read-only extraction of highlighted runs (paragraph index, excerpt, color). Never raises."""
    try:
        from docx.oxml.ns import qn

        highlights: List[Dict[str, Any]] = []
        for para_idx, para in enumerate(doc.paragraphs):
            for run in para.runs:
                rPr = run._element.find(qn("w:rPr"))
                if rPr is None:
                    continue
                hl = rPr.find(qn("w:highlight"))
                if hl is None:
                    continue
                highlights.append({
                    "paragraph_index": para_idx,
                    "excerpt": (run.text or "")[:80],
                    "color": hl.get(qn("w:val")) or "",
                })
                if len(highlights) >= max_entries:
                    return highlights
        return highlights
    except Exception:
        return []


def _resolve_run_for_text(doc: Any, search_text: str) -> Optional[Any]:
    """Finds the first run containing search_text (exact, then case-insensitive)."""
    needle = str(search_text or "")
    if not needle:
        return None
    lowered = needle.lower()
    for para in doc.paragraphs:
        for run in para.runs:
            run_text = run.text or ""
            if needle in run_text:
                return run
    for para in doc.paragraphs:
        for run in para.runs:
            run_text = run.text or ""
            if lowered and lowered in run_text.lower():
                return run
    return None


def _build_add_comment_helper() -> Any:
    """Builds the real-OXML add_comment(doc, paragraph, text, author, run) helper.

    python-docx has no native comment API: this helper injects a genuine
    w:comment into the comments part (created automatically when missing) and
    anchors it with commentRangeStart/commentRangeEnd + a reference run. When
    a run is provided the range wraps that exact run; otherwise the comment is
    anchored at the end of the paragraph. Raises on total OXML failure so the
    recipe sees the real error instead of silently polluting the body text.
    """
    def add_comment(
        doc: Any,
        paragraph: Any,
        text: str,
        author: str = "Assistant",
        run: Optional[Any] = None,
    ) -> None:
        from datetime import datetime, timezone
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        comments_part = _ensure_comments_part(doc)

        existing = comments_part.element.findall(qn("w:comment"))
        comment_id = str(
            max((int(c.get(qn("w:id"), 0)) for c in existing), default=-1) + 1
        )

        initials = "".join(part[0] for part in str(author).split() if part)[:4].upper()
        comment_date = datetime.now(timezone.utc).isoformat()

        comment_el = OxmlElement("w:comment")
        comment_el.set(qn("w:id"), comment_id)
        comment_el.set(qn("w:author"), str(author))
        comment_el.set(qn("w:initials"), initials)
        comment_el.set(qn("w:date"), comment_date)

        comment_p = OxmlElement("w:p")
        comment_r = OxmlElement("w:r")
        comment_t = OxmlElement("w:t")
        comment_t.text = str(text)
        comment_t.set(qn("xml:space"), "preserve")
        comment_r.append(comment_t)
        comment_p.append(comment_r)
        comment_el.append(comment_p)
        comments_part.element.append(comment_el)

        range_start = OxmlElement("w:commentRangeStart")
        range_start.set(qn("w:id"), comment_id)
        range_end = OxmlElement("w:commentRangeEnd")
        range_end.set(qn("w:id"), comment_id)
        reference_run = OxmlElement("w:r")
        comment_ref = OxmlElement("w:commentReference")
        comment_ref.set(qn("w:id"), comment_id)
        reference_run.append(comment_ref)

        if run is not None:
            run._element.addprevious(range_start)
            run._element.addnext(range_end)
            range_end.addnext(reference_run)
        else:
            paragraph._p.append(range_start)
            paragraph._p.append(range_end)
            paragraph._p.append(reference_run)

    return add_comment


def _build_recipe_namespace(working_dir: Optional[Path] = None) -> Dict[str, Any]:
    """Pre-populates the recipe namespace with python-docx design symbols (lazy imports)."""
    namespace: Dict[str, Any] = {}
    try:
        from docx import Document as DocxDocument
        from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
        from docx.enum.text import (
            WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK, WD_COLOR_INDEX,
        )
        from docx.enum.section import WD_SECTION, WD_ORIENT
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        namespace.update({
            "DocxDocument": DocxDocument,
            "DocxInches": DocxInches,
            "DocxPt": DocxPt,
            "DocxRGBColor": DocxRGBColor,
            "WD_ALIGN_PARAGRAPH": WD_ALIGN_PARAGRAPH,
            "WD_LINE_SPACING": WD_LINE_SPACING,
            "WD_BREAK": WD_BREAK,
            "WD_COLOR_INDEX": WD_COLOR_INDEX,
            "WD_SECTION": WD_SECTION,
            "WD_ORIENT": WD_ORIENT,
            "WD_TABLE_ALIGNMENT": WD_TABLE_ALIGNMENT,
            "qn": qn,
            "OxmlElement": OxmlElement,
        })
    except ImportError:
        ASCIIColors.warning("[word_tools] python-docx unavailable — docx recipes disabled.")
        return namespace

    namespace["add_comment"] = _build_add_comment_helper()
    namespace["list_comments"] = lambda doc: _list_comments(doc)
    namespace["list_highlights"] = lambda doc: _list_highlights(doc)
    namespace["resolve_comment_runs"] = _resolve_run_for_text

    if working_dir is not None:
        namespace["media"] = lambda ref: str(_resolve_media_path(ref, working_dir))
        namespace["resolve_path"] = lambda ref: str(_resolve_media_path(ref, working_dir))

        def _list_all_media() -> List[str]:
            aggregated: List[str] = []
            for root in _candidate_working_dirs(working_dir):
                for rel in _list_media_files(root):
                    if rel not in aggregated:
                        aggregated.append(rel)
            return sorted(aggregated, key=str.lower)

        namespace["list_media"] = _list_all_media

    return namespace


def _reopen_office_file(path: Path, suffix: str) -> Tuple[bool, str]:
    """Layer-3 format-aware re-open. Returns (ok, reason). Never raises."""
    try:
        from docx import Document
        doc = Document(path)
        if len(doc.paragraphs) == 0 and len(doc.tables) == 0:
            return False, "re-open succeeded but the document is empty (no paragraphs, no tables)"
        return True, "ok"
    except ImportError:
        return True, "ok (layers 1-2 only; python-docx unavailable)"
    except Exception as reopen_err:
        return False, f"re-open failed: {reopen_err}"


def _validate_office_file(path) -> Tuple[bool, str]:
    """Three-layer Office OOXML integrity validation. Never raises."""
    office_path = Path(path)
    try:
        if not office_path.is_file():
            return False, "file does not exist"
        suffix = office_path.suffix.lower()
        validator = OOXML_VALIDATORS.get(suffix, {})
        required_parts = validator.get("required_parts", ("[Content_Types].xml",))
        with zipfile.ZipFile(office_path) as z:
            bad_member = z.testzip()
            if bad_member:
                return False, f"corrupt zip member: {bad_member}"
            names = z.namelist()
            for required in required_parts:
                if required not in names:
                    return False, f"missing required part: {required}"
            for member_name in names:
                if member_name.endswith((".xml", ".rels")):
                    try:
                        ET.fromstring(z.read(member_name))
                    except ET.ParseError as xml_err:
                        return False, f"malformed XML in part '{member_name}': {xml_err}"
        reopen_ok, reopen_reason = _reopen_office_file(office_path, suffix)
        if not reopen_ok:
            return False, reopen_reason
        return True, reopen_reason
    except zipfile.BadZipFile:
        return False, "not a valid zip container (BadZipFile)"
    except Exception as e:
        return False, f"validation error: {e}"


def _save_and_verify(obj: Any, save_path: Path) -> Dict[str, Any]:
    """Transactional save: snapshot → save → validate → rollback on failure."""
    save_path = Path(save_path)
    pre_existing = save_path.is_file()
    snapshot_bytes = save_path.read_bytes() if pre_existing else None
    try:
        obj.save(str(save_path))
    except Exception as save_err:
        if pre_existing:
            save_path.write_bytes(snapshot_bytes)
        else:
            save_path.unlink(missing_ok=True)
        return {
            "success": False,
            "error": (
                f"Save failed before validation could run: {save_err}\n"
                "The file was restored to its pre-save state. Fix the recipe and retry."
            ),
        }

    ok, reason = _validate_office_file(save_path)
    if ok:
        return {
            "success": True,
            "saved_path": str(save_path),
            "file_name": save_path.name,
        }

    if pre_existing:
        save_path.write_bytes(snapshot_bytes)
        restoration_note = "The file was restored to its pre-save state (pre-save in-memory snapshot)."
    else:
        save_path.unlink(missing_ok=True)
        restoration_note = "The new file was deleted (no previous version to restore)."

    ASCIIColors.error(f"[Word Guard] {save_path.name} failed validation ({reason}) — rolled back.")
    return {
        "success": False,
        "error": (
            f"Office Document Integrity Failure [{OFFICE_GUARD_FAILURE_CODE}]: the saved document "
            f"failed strict validation and was rolled back.\n"
            f"File: {save_path.name}\n"
            f"Reason: {reason}\n"
            f"{restoration_note}\n"
            f"{OFFICE_GUARD_GUIDANCE}"
        ),
    }


def _exec_recipe(recipe: str, namespace: Dict[str, Any], source_name: str) -> Optional[Dict[str, Any]]:
    """Validates and executes a recipe in the sandboxed namespace."""
    if not recipe or not recipe.strip():
        return None
    if len(recipe) > _MAX_RECIPE_CHARS:
        return {
            "success": False,
            "error": f"Recipe exceeds the maximum length of {_MAX_RECIPE_CHARS} characters.",
        }
    try:
        _check_recipe(recipe)
    except ValueError as forbidden_err:
        return {"success": False, "error": str(forbidden_err)}

    dedented = textwrap.dedent(recipe)
    namespace["__builtins__"] = _build_safe_builtins()
    try:
        exec(compile(dedented, source_name, "exec"), namespace)
        return None
    except Exception:
        annotated = _annotate_traceback(traceback.format_exc(), dedented, source_name)
        if "save(" in annotated and "missing" in annotated:
            annotated += (
                "\nHint: the tool saves the file automatically after your recipe — "
                "remove the save() call."
            )
        return {"success": False, "error": f"Recipe execution error:\n{annotated}"}


def _resolve_working_dir(discussion_instance: Optional[Any]) -> Path:
    """Resolves the directory files are saved to / read from (discussion workspace or CWD)."""
    if discussion_instance is not None:
        workspace = getattr(discussion_instance, "workspace_data_path", None)
        if workspace:
            return Path(workspace)
    return Path(".")


def _sanitize_file_name(file: str, default: str) -> str:
    """Normalizes a file name, strips path components, and applies the default when empty."""
    name = (file or "").strip()
    if not name:
        return default
    name = Path(name).name
    if not name or name in (".", ".."):
        return default
    return name


def _create_root_object(suffix: str, namespace: Dict[str, Any]) -> Any:
    """Creates a fresh document and binds it into the namespace as doc."""
    if suffix in (".docx", ".docm"):
        doc = namespace["DocxDocument"]()
        namespace["doc"] = doc
        return doc
    raise ValueError(f"Unsupported Word file extension: {suffix}")


def _open_existing(path: Path, namespace: Dict[str, Any]) -> Any:
    """Opens an existing document by suffix and binds it into the namespace."""
    suffix = path.suffix.lower()
    if suffix in (".docx", ".docm"):
        doc = namespace["DocxDocument"](str(path))
        namespace["doc"] = doc
        return doc
    raise ValueError(f"Unsupported Word file extension: {suffix}")


def _pick_effective_root(namespace: Dict[str, Any], suffix: str, original_root: Any) -> Any:
    """Returns the recipe-rebound root alias when the recipe replaced the
    pre-bound root object (e.g. `doc = DocxDocument()`), else the original root."""
    aliases = {".docx": ("doc",), ".docm": ("doc",)}
    for alias in aliases.get(suffix, ()):
        candidate = namespace.get(alias)
        if candidate is not None and candidate is not original_root:
            return candidate
    return original_root


def _build_result(save_result: Dict[str, Any], save_path: Path) -> Dict[str, Any]:
    """Shapes a successful save result into the standard result contract."""
    return {
        "success": True,
        "workspace_file": str(save_path.resolve()),
        "context_files": [str(save_path.resolve())],
        "file_name": save_path.name,
    }


def _require_suffix(
    file: str,
    allowed_suffixes: Tuple[str, ...],
    tool_name: str,
) -> Optional[Dict[str, Any]]:
    """Returns a structured failure dict when the file suffix is outside the
    tool's format family, else None (proceed)."""
    actual = Path(str(file or "")).suffix.lower()
    if actual not in allowed_suffixes:
        return {
            "success": False,
            "error": (
                f"{tool_name} requires a {'/'.join(allowed_suffixes)} file, "
                f"got '{actual or '(no extension)'}'. Use the tool matching "
                f"this file's format instead."
            ),
            "code": "wrong_file_type_for_tool",
        }
    return None


def _read_docx(path: Path) -> str:
    """Builds a markdown inventory of a document (paragraphs, tables, comments, highlights)."""
    from docx import Document

    doc = Document(str(path))
    lines = [f"# Document: {path.name}", ""]
    paragraph_lines = []
    for para_idx, para in enumerate(doc.paragraphs):
        if para.text.strip():
            paragraph_lines.append(f"- [{para_idx}] ({para.style.name}): {para.text}")
    lines.append("## Paragraphs")
    lines.extend(paragraph_lines if paragraph_lines else ["- (no paragraphs)"])
    lines.append("")
    lines.append(f"## Tables: {len(doc.tables)}")
    for table_idx, table in enumerate(doc.tables, start=1):
        lines.append(f"- table {table_idx}: {len(table.rows)} rows x {len(table.columns)} cols")

    lines.append("")
    lines.append("## Comments")
    comments = _list_comments(doc)
    if comments:
        for i, comment in enumerate(comments, start=1):
            lines.append(
                f"- [{i}] {comment['author']} ({comment['date']}): {comment['text']}"
            )
    else:
        lines.append("- (no comments)")

    lines.append("")
    lines.append("## Highlights")
    highlights = _list_highlights(doc)
    if highlights:
        for hl in highlights:
            lines.append(
                f"- paragraph {hl['paragraph_index']}: \"{hl['excerpt']}\" ({hl['color']})"
            )
    else:
        lines.append("- (none)")

    return "\n".join(lines)


def tool_build_word(
    file: str = "document.docx",
    recipe: str = "",
    discussion_instance: Optional[Any] = None,
    lollms_client_instance: Optional[Any] = None,
) -> Dict[str, Any]:
    """Creates a new Word document (.docx/.docm) from a sandboxed python-docx recipe.

    === WORD DOCTRINE (HINTS — YOU CODE FREELY) ===
    The root object is pre-bound as `doc`. Available symbols: DocxDocument,
    DocxInches, DocxPt, DocxRGBColor, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING,
    WD_BREAK, WD_COLOR_INDEX, WD_SECTION, WD_ORIENT, WD_TABLE_ALIGNMENT, qn,
    OxmlElement, add_comment, list_comments, list_highlights,
    resolve_comment_runs, media(), list_media(), resolve_path().
    WORKSPACE RULE: every workspace asset reference (images, sub-documents,
    inputs) MUST go through media()/resolve_path() — bare filenames resolve
    against the workspace, not the process folder.

    Building hints (not restrictions — use your own coding skills):
    - Rich text: doc.add_paragraph("text", style="Heading 1"), runs with
      run.bold/italic/underline, run.font.color.rgb = DocxRGBColor(0x1F,0x4E,0x79).
    - Styles: doc.styles, paragraph.style = doc.styles["Title"].
    - Tables: doc.add_table(rows, cols), cell merging, WD_TABLE_ALIGNMENT.
    - Images: doc.add_picture(media("relative/path.png"), width=DocxInches(4)).
    - Headers/footers: section.header / section.footer paragraphs.
    - TOC: insert a TOC field via OXML (fldSimple with qn) — Word populates
      it on open (user presses F9 or enables auto-update).
    - Sections: doc.add_section(WD_SECTION.NEW_PAGE) for multi-part documents.

    Recipes are plain Python: imports are allowed ONLY from the whitelist
    (docx.* and safe stdlib modules); all core design symbols are also
    pre-imported. The tool saves automatically; calling doc.save() inside the
    recipe is a no-op — remove it.

    Args:
        file: Target .docx/.docm file name (default "document.docx").
        recipe: Plain Python recipe code mutating the pre-bound doc object.
        discussion_instance: Active discussion session instance.
        lollms_client_instance: Active client instance.
    """
    suffix_failure = _require_suffix(file, _DOCX_SUFFIXES, "tool_build_word")
    if suffix_failure is not None:
        return suffix_failure
    try:
        file_name = _sanitize_file_name(file, "document.docx")
        suffix = Path(file_name).suffix.lower()
        if suffix not in GUARDED_EXTENSIONS:
            return {
                "success": False,
                "error": f"Unsupported Word file extension '{suffix}'. Supported: {', '.join(GUARDED_EXTENSIONS)}.",
            }

        working_dir = _resolve_working_dir(discussion_instance)
        namespace = _build_recipe_namespace(working_dir=working_dir)
        try:
            root_obj = _create_root_object(suffix, namespace)
        except KeyError as missing_symbol:
            return {
                "success": False,
                "error": (
                    f"python-docx is not installed (missing symbol: {missing_symbol}). "
                    "Ask the user to install it and retry."
                ),
            }

        try:
            root_obj.save = lambda *args, **kwargs: None  # the tool performs the real validated save
        except Exception:
            pass

        exec_failure = _exec_recipe(recipe, namespace, "<word_recipe>")
        if exec_failure is not None:
            return exec_failure

        root_obj = _pick_effective_root(namespace, suffix, root_obj)

        save_path = working_dir / file_name
        save_result = _save_and_verify(root_obj, save_path)
        if not save_result["success"]:
            return save_result
        return _build_result(save_result, save_path)
    except Exception as unexpected_err:
        return {
            "success": False,
            "error": f"Unexpected failure building Word document: {unexpected_err}",
        }


def tool_edit_word(
    file: str,
    recipe: str = "",
    discussion_instance: Optional[Any] = None,
    lollms_client_instance: Optional[Any] = None,
) -> Dict[str, Any]:
    """Edits an existing Word document (.docx/.docm) in place via a sandboxed python-docx recipe.

    === WORD EDIT DOCTRINE (HINTS — YOU CODE FREELY) ===
    The opened document is pre-bound as `doc`. Available symbols: DocxDocument,
    DocxInches, DocxPt, DocxRGBColor, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING,
    WD_BREAK, WD_COLOR_INDEX, WD_SECTION, WD_ORIENT, WD_TABLE_ALIGNMENT, qn,
    OxmlElement, add_comment, list_comments, list_highlights,
    resolve_comment_runs, media(), list_media(), resolve_path().
    WORKSPACE RULE: every workspace asset reference (images, sub-documents,
    inputs) MUST go through media()/resolve_path() — bare filenames resolve
    against the workspace, not the process folder.

    Edit vocabulary hints (not restrictions — use your own coding skills):
    - INSERT COMMENTS: add_comment(doc, paragraph, "comment text", author="Assistant",
      run=resolve_comment_runs(doc, "target text")) — real OXML comments
      (w:comment in comments.xml, anchored with commentRangeStart/End around
      the exact run). The comments part is created automatically when missing;
      resolve_comment_runs(doc, "text") returns the first run containing the
      text so the comment anchors precisely on it.
    - REVIEW STATE: list_comments(doc) and list_highlights(doc) return the
      existing comments (author/date/text) and highlighted runs.
    - HIGHLIGHT TEXT: run.font.highlight_color = WD_COLOR_INDEX.YELLOW
      (also GREEN, PINK, CYAN, RED...).
    - REFORMAT: paragraph.style = doc.styles["Heading 2"],
      run.font.size = DocxPt(14), run.font.name = "Calibri",
      paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER,
      paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE,
      run.font.color.rgb = DocxRGBColor(0xC0, 0x00, 0x00).
    - ADD SECTIONS: doc.add_section(WD_SECTION.NEW_PAGE) (also CONTINUOUS,
      ODD_PAGE; set orientation via section.orientation = WD_ORIENT.LANDSCAPE).
    - UPDATE LINKS / FIELDS: python-docx cannot refresh external links or
      field results. Insert field CODES via OXML so Word updates them on
      open: build a w:fldSimple (or fldChar run sequence) with qn("w:instrText")
      containing e.g. 'TOC \\o "1-3" \\h \\z \\u' or 'HYPERLINK "url"'.
    - FIND & REPLACE: iterate doc.paragraphs (and table cells), match
      paragraph.runs text spans, rewrite run.text.
    - TABLES: doc.tables[i], cell(r,c), add rows, merge cells.
    - IMAGES: paragraph.add_run().add_picture(media("img.png"), width=DocxInches(3)).

    Recipes are plain Python: imports are allowed ONLY from the whitelist
    (docx.* and safe stdlib modules); all core design symbols are also
    pre-imported. The tool saves automatically; calling doc.save() inside the
    recipe is a no-op — remove it. The save is transactional: on validation
    failure the pre-edit bytes are restored.

    Args:
        file: Existing .docx/.docm file to edit (relative to workspace or absolute path).
        recipe: Plain Python recipe code mutating the opened doc object.
        discussion_instance: Active discussion session instance.
        lollms_client_instance: Active client instance.
    """
    suffix_failure = _require_suffix(file, _DOCX_SUFFIXES, "tool_edit_word")
    if suffix_failure is not None:
        return suffix_failure
    try:
        if not file or not str(file).strip():
            return {"success": False, "error": "No file provided for editing."}

        working_dir = _resolve_working_dir(discussion_instance)
        candidate = Path(str(file).strip())
        file_path = candidate if candidate.is_absolute() else working_dir / candidate
        if not file_path.is_file():
            return {"success": False, "error": f"File not found: {file_path}"}

        suffix = file_path.suffix.lower()
        if suffix not in GUARDED_EXTENSIONS:
            return {
                "success": False,
                "error": f"Unsupported Word file extension '{suffix}'. Supported: {', '.join(GUARDED_EXTENSIONS)}.",
            }

        namespace = _build_recipe_namespace(working_dir=working_dir)
        try:
            root_obj = _open_existing(file_path, namespace)
        except KeyError as missing_symbol:
            return {
                "success": False,
                "error": (
                    f"python-docx is not installed (missing symbol: {missing_symbol}). "
                    "Ask the user to install it and retry."
                ),
            }

        try:
            root_obj.save = lambda *args, **kwargs: None  # the tool performs the real validated save
        except Exception:
            pass

        exec_failure = _exec_recipe(recipe, namespace, "<word_recipe>")
        if exec_failure is not None:
            return exec_failure

        root_obj = _pick_effective_root(namespace, suffix, root_obj)

        save_result = _save_and_verify(root_obj, file_path)
        if not save_result["success"]:
            return save_result
        return _build_result(save_result, file_path)
    except Exception as unexpected_err:
        return {
            "success": False,
            "error": f"Unexpected failure editing Word document: {unexpected_err}",
        }


def tool_read_word(
    file: str,
    discussion_instance: Optional[Any] = None,
    lollms_client_instance: Optional[Any] = None,
) -> Dict[str, Any]:
    """Reads an existing Word document (.docx/.docm) and returns a markdown inventory.

    Returns paragraphs (with style names) and tables. Validates the file
    first; corrupt files return a structured failure instead of a crash.
    No recipe is executed and nothing is saved.
    """
    try:
        if not file or not str(file).strip():
            return {"success": False, "error": "No file provided for reading."}

        working_dir = _resolve_working_dir(discussion_instance)
        candidate = Path(str(file).strip())
        file_path = candidate if candidate.is_absolute() else working_dir / candidate
        if not file_path.is_file():
            return {"success": False, "error": f"File not found: {file_path}"}

        suffix = file_path.suffix.lower()
        if suffix not in _DOCX_SUFFIXES:
            return {
                "success": False,
                "error": f"tool_read_word requires a {'/'.join(_DOCX_SUFFIXES)} file, got '{suffix or '(no extension)'}'.",
            }

        ok, reason = _validate_office_file(file_path)
        if not ok:
            return {
                "success": False,
                "error": (
                    f"Office Document Integrity Failure [{OFFICE_GUARD_FAILURE_CODE}]: "
                    f"'{file_path.name}' is corrupt and cannot be read safely.\n"
                    f"Reason: {reason}"
                ),
            }

        inventory = _read_docx(file_path)
        return {
            "success": True,
            "file_name": file_path.name,
            "workspace_file": str(file_path.resolve()),
            "inventory": inventory,
        }
    except Exception as unexpected_err:
        return {
            "success": False,
            "error": f"Unexpected failure reading Word document: {unexpected_err}",
        }


def tool_word_review(
    file: str,
    discussion_instance: Optional[Any] = None,
    lollms_client_instance: Optional[Any] = None,
) -> Dict[str, Any]:
    """Reads the review state of a Word document (.docx/.docm): comments and highlights.

    Returns the structured list of comments (author, initials, date, text) and
    highlighted runs (paragraph index, excerpt, color), plus the full markdown
    inventory. Read-only: no recipe is executed and nothing is saved. Use this
    to verify annotations you added or to review feedback before editing.
    """
    try:
        if not file or not str(file).strip():
            return {"success": False, "error": "No file provided for review."}

        working_dir = _resolve_working_dir(discussion_instance)
        candidate = Path(str(file).strip())
        file_path = candidate if candidate.is_absolute() else working_dir / candidate
        if not file_path.is_file():
            return {"success": False, "error": f"File not found: {file_path}"}

        suffix = file_path.suffix.lower()
        if suffix not in _DOCX_SUFFIXES:
            return {
                "success": False,
                "error": f"tool_word_review requires a {'/'.join(_DOCX_SUFFIXES)} file, got '{suffix or '(no extension)'}'.",
            }

        ok, reason = _validate_office_file(file_path)
        if not ok:
            return {
                "success": False,
                "error": (
                    f"Office Document Integrity Failure [{OFFICE_GUARD_FAILURE_CODE}]: "
                    f"'{file_path.name}' is corrupt and cannot be read safely.\n"
                    f"Reason: {reason}"
                ),
            }

        from docx import Document
        doc = Document(str(file_path))
        comments = _list_comments(doc)
        highlights = _list_highlights(doc)
        inventory = _read_docx(file_path)

        return {
            "success": True,
            "file_name": file_path.name,
            "workspace_file": str(file_path.resolve()),
            "comments": comments,
            "highlights": highlights,
            "inventory": inventory,
        }
    except Exception as unexpected_err:
        return {
            "success": False,
            "error": f"Unexpected failure reviewing Word document: {unexpected_err}",
        }


def get_tools() -> List[Dict[str, Any]]:
    """Returns the tool specifications for the Word Tools toolset."""
    return [
        {
            "name": "tool_build_word",
            "description": (
                "Creates a new Word document (.docx/.docm) from a sandboxed python-docx recipe. "
                "Root object pre-bound as doc. Symbols: DocxDocument, DocxInches, DocxPt, "
                "DocxRGBColor, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK, WD_COLOR_INDEX, "
                "WD_SECTION, WD_ORIENT, WD_TABLE_ALIGNMENT, qn, OxmlElement, add_comment, "
                "list_comments, list_highlights, resolve_comment_runs, media(), list_media(), "
                "resolve_path(). WORKSPACE RULE: every workspace asset reference (images, "
                "inputs) MUST use media()/resolve_path() — bare filenames resolve against "
                "the workspace, not the process folder. "
                "Hints: styles, tables, images via media(), headers/footers, TOC field "
                "via OXML, sections via doc.add_section(WD_SECTION.NEW_PAGE). Comments are "
                "REAL Word comments (comments part auto-created, run-anchored). Imports "
                "allowed ONLY from the whitelist (docx.*, safe stdlib). Do NOT call save() "
                "— the tool saves and validates automatically."
            ),
            "parameters": [
                {
                    "name": "file",
                    "type": "str",
                    "description": "Target .docx/.docm file name (default 'document.docx').",
                    "optional": True,
                },
                {
                    "name": "recipe",
                    "type": "str",
                    "description": "Plain Python recipe mutating the pre-bound doc. Imports allowed only from the whitelist (docx.*, safe stdlib). Do NOT call save().",
                    "optional": True,
                },
            ],
            "output": [
                {"name": "success", "type": "bool", "description": "Whether the operation succeeded."},
                {"name": "workspace_file", "type": "str", "description": "Absolute path to the saved file."},
                {"name": "context_files", "type": "list", "description": "List of context files (the saved file)."},
                {"name": "file_name", "type": "str", "description": "Name of the saved file."},
                {"name": "error", "type": "str", "description": "Error message if failed."},
            ],
            "callable": tool_build_word,
            "_source": "word_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_edit_word",
            "description": (
                "Edits an existing Word document (.docx/.docm) in place via a sandboxed python-docx "
                "recipe. Opened document pre-bound as doc. EDIT VOCABULARY (hints, not restrictions): "
                "REAL comments via add_comment(doc, paragraph, text, author, run=resolve_comment_runs("
                "doc, 'target text')) — the comments part is auto-created and the comment anchors on "
                "the exact run; review state via list_comments(doc)/list_highlights(doc); highlights "
                "via run.font.highlight_color = WD_COLOR_INDEX.YELLOW; reformat via styles/fonts/"
                "alignment/spacing; sections via doc.add_section(WD_SECTION.NEW_PAGE); links/fields "
                "via OXML field codes (qn, fldSimple — Word updates on open); find & replace over "
                "runs; tables; images via media(). WORKSPACE RULE: every workspace asset "
                "reference (images, inputs) MUST use media()/resolve_path() — bare filenames "
                "resolve against the workspace, not the process folder. Imports allowed ONLY from the whitelist (docx.*, "
                "safe stdlib). Do NOT call save() — the tool saves and validates automatically. "
                "Transactional save: pre-edit bytes restored on validation failure."
            ),
            "parameters": [
                {
                    "name": "file",
                    "type": "str",
                    "description": "Existing .docx/.docm file to edit (relative to workspace or absolute path).",
                },
                {
                    "name": "recipe",
                    "type": "str",
                    "description": "Plain Python recipe mutating the opened doc. Comments: add_comment(doc, paragraph, text, author, run=resolve_comment_runs(doc, 'text')). Highlights: WD_COLOR_INDEX. Review: list_comments/list_highlights. Sections: WD_SECTION. Fields: OXML via qn. Imports allowed only from the whitelist (docx.*, safe stdlib). Do NOT call save().",
                    "optional": True,
                },
            ],
            "output": [
                {"name": "success", "type": "bool", "description": "Whether the operation succeeded."},
                {"name": "workspace_file", "type": "str", "description": "Absolute path to the saved file."},
                {"name": "context_files", "type": "list", "description": "List of context files (the saved file)."},
                {"name": "file_name", "type": "str", "description": "Name of the saved file."},
                {"name": "error", "type": "str", "description": "Error message if failed."},
            ],
            "callable": tool_edit_word,
            "_source": "word_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_read_word",
            "description": (
                "Reads an existing Word document (.docx/.docm) and returns a markdown inventory "
                "(paragraphs with styles, tables, comments, highlights). Validates the file first; "
                "corrupt files return a structured failure."
            ),
            "parameters": [
                {
                    "name": "file",
                    "type": "str",
                    "description": "Existing .docx/.docm file to read (relative to workspace or absolute path).",
                },
            ],
            "output": [
                {"name": "success", "type": "bool", "description": "Whether the operation succeeded."},
                {"name": "inventory", "type": "str", "description": "Markdown inventory of the document contents."},
                {"name": "file_name", "type": "str", "description": "Name of the file read."},
                {"name": "workspace_file", "type": "str", "description": "Absolute path to the file."},
                {"name": "error", "type": "str", "description": "Error message if failed."},
            ],
            "callable": tool_read_word,
            "_source": "word_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_word_review",
            "description": (
                "Reads the review state of an existing Word document (.docx/.docm): returns the "
                "structured list of comments (author, initials, date, text) and highlighted runs "
                "(paragraph index, excerpt, color), plus the full markdown inventory. Read-only: "
                "no recipe is executed and nothing is saved. Use it to verify annotations you "
                "added or to review feedback before editing."
            ),
            "parameters": [
                {
                    "name": "file",
                    "type": "str",
                    "description": "Existing .docx/.docm file to review (relative to workspace or absolute path).",
                },
            ],
            "output": [
                {"name": "success", "type": "bool", "description": "Whether the operation succeeded."},
                {"name": "comments", "type": "list", "description": "Structured comments (author, initials, date, text)."},
                {"name": "highlights", "type": "list", "description": "Highlighted runs (paragraph index, excerpt, color)."},
                {"name": "inventory", "type": "str", "description": "Full markdown inventory of the document."},
                {"name": "file_name", "type": "str", "description": "Name of the file reviewed."},
                {"name": "workspace_file", "type": "str", "description": "Absolute path to the file."},
                {"name": "error", "type": "str", "description": "Error message if failed."},
            ],
            "callable": tool_word_review,
            "_source": "word_tools",
            "_source_file": __file__,
        },
    ]