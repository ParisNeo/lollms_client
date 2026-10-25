"""
Excel Tools — dedicated LCP toolset for Microsoft Excel workbooks.

Doctrine (read this before writing any recipe):
- Excel (.xlsx) is a BINARY OOXML workbook (a ZIP of XML parts). NEVER write
  plain text or markdown into an .xlsx file — it becomes unreadable. Build
  workbooks ONLY via these tools: openpyxl recipes (tool_build_excel /
  tool_edit_excel) or tool_excel_from_markdown for markdown/structured-text
  input.
- NEVER write computed results manually. ALWAYS write FORMULAS into cells so
  Excel computes them: ws["F2"] = "=SUM(B2:E2)", ws["G2"] = "=AVERAGE(B2:E2)",
  ws["H2"] = '=IF(F2>100,"High","Low")'. Use SUM, AVERAGE, MIN, MAX, COUNT,
  IF, VLOOKUP, INDEX/MATCH, CONCATENATE, etc. Literal values are ONLY for
  input data.
- Chart caveat: when a chart's Reference points at formula cells, openpyxl
  writes formulas WITHOUT cached values — some desktop viewers show an empty
  chart until Excel recalculates. Formulas remain the preferred approach
  (validation passes structurally); precomputed literals are acceptable only
  when the user needs chart rendering without recalculation.
"""
import re
import textwrap
import traceback
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from ascii_colors import ASCIIColors

TOOL_LIBRARY_NAME = "Excel Tools"
TOOL_LIBRARY_DESC = (
    "Create, edit and read Excel workbooks (.xlsx/.xlsm) via sandboxed openpyxl "
    "recipe code, or convert markdown documents into real workbooks. DOCTRINE: "
    "Excel (.xlsx) is a binary OOXML workbook (ZIP of XML parts) — NEVER paste "
    "markdown or plain text into an .xlsx file, it becomes unreadable; produce "
    "workbooks ONLY through these tools (openpyxl recipes or "
    "tool_excel_from_markdown). Never write computed results manually — always "
    "write formulas (=SUM, =AVERAGE, =IF, =VLOOKUP ...) so Excel computes them. "
    "Literal values only for input data. Every save is verified with a strict "
    "three-layer OOXML integrity check before being committed to the workspace."
)
TOOL_LIBRARY_ICON = "📈"

_MAX_RECIPE_CHARS = 60_000
_FORBIDDEN_TOKENS = (
    "eval(", "exec(", "compile(", "open(",
    "os.", "sys.", "subprocess", "shutil", "pathlib", "__",
    "globals(", "locals(", "vars(", "getattr(", "setattr(", "delattr(",
)

_IMPORT_PREFIX_WHITELIST = ("openpyxl.",)
_SAFE_STDLIB_MODULES = frozenset({
    "math", "datetime", "random", "json", "statistics", "itertools",
    "functools", "re", "uuid", "calendar", "decimal", "fractions",
    "textwrap", "base64", "hashlib", "string", "collections",
})

_MEDIA_EXTENSIONS = (
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg", ".tiff", ".ico",
    ".mp4", ".mov", ".avi", ".mkv", ".webm", ".wmv",
)
_MAX_MEDIA_SCAN = 5000
_MAX_MEDIA_LIST = 2000

OOXML_VALIDATORS = {
    ".xlsx": {"required_parts": ("[Content_Types].xml", "xl/workbook.xml", "xl/_rels/workbook.xml.rels"), "reopener": "openpyxl"},
    ".xlsm": {"required_parts": ("[Content_Types].xml", "xl/workbook.xml", "xl/_rels/workbook.xml.rels"), "reopener": "openpyxl"},
}
GUARDED_EXTENSIONS = tuple(OOXML_VALIDATORS.keys())

OFFICE_GUARD_FAILURE_CODE = "office_file_failed_validation"

OFFICE_GUARD_GUIDANCE = (
    "Likely causes: chart anchored over data/helper ranges, chart width/height mutated "
    "after add_chart(), dropped native parts (richData/VBA/slicers) when re-saving a "
    "native workbook, or a structural defect in the produced workbook. Do NOT re-issue "
    "identical code. Rebuild the content: set chart width/height BEFORE add_chart(), "
    "anchor charts away from data/helper ranges, and prefer building fresh workbooks "
    "over round-tripping desktop-authored files."
)

_XLSX_SUFFIXES = (".xlsx", ".xlsm")


def init_tools_library(config: dict = None) -> None:
    try:
        import pipmaster as pm
        pm.ensure_packages(["openpyxl"])
    except Exception as e:
        ASCIIColors.warning(f"[excel_tools] Failed to ensure dependencies: {e}")


def _safe_import(name: Any, *args: Any, **kwargs: Any) -> Any:
    """Whitelist gate for recipe imports (installed as __import__ in the sandbox)."""
    module_name = str(name)
    root_mod = module_name.split(".")[0]
    if root_mod in _SAFE_STDLIB_MODULES or any(
        module_name == prefix.rstrip(".") or module_name.startswith(prefix)
        for prefix in _IMPORT_PREFIX_WHITELIST
    ):
        import importlib
        return importlib.import_module(module_name)
    raise ImportError(
        f"Import '{module_name}' is not allowed in recipes. Allowed imports: openpyxl.* "
        "and safe stdlib modules (math, datetime, random, json, statistics, itertools, "
        "functools, re, uuid, calendar, decimal, fractions, textwrap, base64, hashlib, "
        "string, collections). All design symbols are also pre-imported."
    )


def _build_safe_builtins() -> Dict[str, Any]:
    """Curated builtins for the recipe sandbox: no file/system/eval access."""
    import builtins
    base = {k: v for k, v in vars(builtins).items() if not k.startswith("_")}
    for dangerous in ("open", "eval", "exec", "compile", "breakpoint", "input",
                      "exit", "quit", "help", "license", "globals", "locals",
                      "vars", "getattr", "setattr", "delattr"):
        base.pop(dangerous, None)
    base["__import__"] = _safe_import
    return base


def _check_recipe(recipe: str) -> None:
    """Rejects recipes containing forbidden tokens before execution."""
    for token in _FORBIDDEN_TOKENS:
        if token in recipe:
            raise ValueError(
                f"Forbidden token '{token}' in recipe. Recipes are plain Python: imports "
                "are allowed only from the whitelist (openpyxl.* and safe stdlib modules), "
                "with no file/system access and no dunder access. Prefer the pre-imported "
                "design symbols."
            )


def _annotate_traceback(raw_traceback: str, source_code: str, source_name: str) -> str:
    """Appends the exact failing recipe line after each frame referencing source_name."""
    code_lines = source_code.splitlines()
    annotated_lines = []
    for line in raw_traceback.splitlines():
        annotated_lines.append(line)
        match = re.search(
            re.escape(f'File "{source_name}", line ') + r"(\d+)", line
        )
        if match:
            failing_line_no = int(match.group(1))
            if 1 <= failing_line_no <= len(code_lines):
                annotated_lines.append(
                    f"    >>> LINE {failing_line_no}: {code_lines[failing_line_no - 1]}"
                )
    return "\n".join(annotated_lines)


def _candidate_working_dirs(working_dir: Optional[Path]) -> List[Path]:
    """Enumerates the roots a media reference may be anchored to."""
    roots: List[Path] = []
    for base in (working_dir, Path(".").resolve()):
        if base is None:
            continue
        try:
            resolved = Path(base).resolve()
        except Exception:
            continue
        if resolved.is_dir() and resolved not in roots:
            roots.append(resolved)
    return roots


def _resolve_media_path(ref: str, working_dir: Optional[Path]) -> Path:
    """Resolves a workspace-relative media reference to an absolute path (read-only)."""
    raw = str(ref or "").strip().strip('"').strip("'")
    if not raw:
        raise ValueError("media() received an empty reference.")
    if "\\" in raw:
        raw = raw.replace("\\", "/")
    while raw.startswith("./"):
        raw = raw[2:]
    if Path(raw).is_absolute() or raw.startswith("/"):
        raise ValueError(f"Media reference '{ref}' must be workspace-relative, not absolute.")
    if ".." in Path(raw).parts:
        raise ValueError(f"Media reference '{ref}' must not contain '..' (workspace escape rejected).")

    roots = _candidate_working_dirs(working_dir)
    if not roots:
        raise ValueError(f"Media file '{ref}' could not be resolved: no accessible workspace directory.")

    target_name = Path(raw).name
    searched_roots: List[str] = []
    for root in roots:
        searched_roots.append(str(root))
        candidate = (root / raw).resolve()
        if candidate.is_relative_to(root) and candidate.is_file():
            return candidate
    for root in roots:
        scanned = 0
        for hit in root.rglob(target_name):
            scanned += 1
            if scanned > _MAX_MEDIA_SCAN:
                break
            hit_resolved = hit.resolve()
            if hit_resolved.is_relative_to(root) and hit_resolved.is_file():
                return hit_resolved
    available = []
    for root in roots:
        available.extend(_list_media_files(root))
        if len(available) >= 20:
            break
    available_hint = ", ".join(available[:20]) if available else "(none found)"
    raise ValueError(
        f"Media file '{ref}' not found. Searched roots: {', '.join(searched_roots)} "
        f"(exact path and recursive name search, subfolders included). "
        f"Available media: {available_hint}. Use list_media() for the full list."
    )


def _list_media_files(working_dir: Path) -> List[str]:
    """Enumerates media files under the workspace (subfolders included). Never raises."""
    root = Path(working_dir).resolve()
    if not root.is_dir():
        return []
    found: List[str] = []
    scanned = 0
    for entry in root.rglob("*"):
        scanned += 1
        if scanned > _MAX_MEDIA_SCAN:
            break
        if entry.suffix.lower() in _MEDIA_EXTENSIONS:
            try:
                rel = entry.resolve().relative_to(root)
            except ValueError:
                continue
            found.append(rel.as_posix())
            if len(found) >= _MAX_MEDIA_LIST:
                break
    return sorted(found, key=str.lower)


def _build_recipe_namespace(working_dir: Optional[Path] = None) -> Dict[str, Any]:
    """Pre-populates the recipe namespace with openpyxl design symbols (lazy imports).

    The `Workbook` symbol is a SHIM FACTORY: every workbook a recipe constructs
    (rebound or not) arrives with a no-op instance-level `save`, so in-recipe
    `wb.save()` calls are absorbed instead of raising TypeError. The real,
    validated save is performed by `_save_and_verify` via the class-level
    genuine method, which bypasses the instance attribute.
    """
    namespace: Dict[str, Any] = {}
    try:
        from openpyxl import Workbook as _RealWorkbook, load_workbook
        from openpyxl.chart import PieChart, LineChart, BarChart, Reference
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
        from openpyxl.worksheet.datavalidation import DataValidation

        def _make_shimmed_workbook(*args: Any, **kwargs: Any) -> Any:
            wb = _RealWorkbook(*args, **kwargs)
            try:
                wb.save = lambda *a, **k: None  # the tool performs the real validated save
            except Exception:
                pass
            return wb

        namespace.update({
            "Workbook": _make_shimmed_workbook,
            "load_workbook": load_workbook,
            "PieChart": PieChart,
            "LineChart": LineChart,
            "BarChart": BarChart,
            "Reference": Reference,
            "Font": Font,
            "PatternFill": PatternFill,
            "Alignment": Alignment,
            "Border": Border,
            "Side": Side,
            "get_column_letter": get_column_letter,
            "CellIsRule": CellIsRule,
            "ColorScaleRule": ColorScaleRule,
            "DataValidation": DataValidation,
        })
    except ImportError:
        ASCIIColors.warning("[excel_tools] openpyxl unavailable — xlsx recipes disabled.")

    if working_dir is not None:
        namespace["media"] = lambda ref: str(_resolve_media_path(ref, working_dir))

        def _list_all_media() -> List[str]:
            aggregated: List[str] = []
            for root in _candidate_working_dirs(working_dir):
                for rel in _list_media_files(root):
                    if rel not in aggregated:
                        aggregated.append(rel)
            return sorted(aggregated, key=str.lower)

        namespace["list_media"] = _list_all_media

    return namespace


def _reopen_office_file(path: Path, suffix: str) -> Tuple[bool, str]:
    """Layer-3 format-aware re-open. Returns (ok, reason). Never raises.

    Structural re-open only: formula cells pass validation (no data_only
    value requirement) — Excel computes formulas when the user opens the file.
    """
    try:
        from openpyxl import load_workbook
        if suffix == ".xlsm":
            wb = load_workbook(path, keep_vba=True)
        else:
            wb = load_workbook(path)
        # STRUCTURAL checks only: formula cells carry no cached values under
        # openpyxl authoring, so demanding data would false-negative every
        # doctrine-compliant workbook and trigger a spurious rollback.
        if not wb.sheetnames:
            return False, "re-open failed: workbook has no sheets"
        if all(wb[name].max_row < 1 for name in wb.sheetnames):
            return False, "re-open failed: all sheets are empty"
        return True, "ok"
    except ImportError:
        return True, "ok (layers 1-2 only; openpyxl unavailable)"
    except Exception as reopen_err:
        return False, f"re-open failed: {reopen_err}"


def _validate_office_file(path) -> Tuple[bool, str]:
    """Three-layer Office OOXML integrity validation. Never raises."""
    office_path = Path(path)
    try:
        if not office_path.is_file():
            return False, "file does not exist"
        suffix = office_path.suffix.lower()
        validator = OOXML_VALIDATORS.get(suffix, {})
        required_parts = validator.get("required_parts", ("[Content_Types].xml",))
        with zipfile.ZipFile(office_path) as z:
            bad_member = z.testzip()
            if bad_member:
                return False, f"corrupt zip member: {bad_member}"
            names = z.namelist()
            for required in required_parts:
                if required not in names:
                    return False, f"missing required part: {required}"
            for member_name in names:
                if member_name.endswith((".xml", ".rels")):
                    try:
                        ET.fromstring(z.read(member_name))
                    except ET.ParseError as xml_err:
                        return False, f"malformed XML in part '{member_name}': {xml_err}"
        reopen_ok, reopen_reason = _reopen_office_file(office_path, suffix)
        if not reopen_ok:
            return False, reopen_reason
        return True, reopen_reason
    except zipfile.BadZipFile:
        return False, "not a valid zip container (BadZipFile)"
    except Exception as e:
        return False, f"validation error: {e}"


def _save_and_verify(obj: Any, save_path: Path) -> Dict[str, Any]:
    """Transactional save: snapshot → save → validate → rollback on failure.

    The recipe sandbox installs a no-op `save` shim on the root object so
    in-recipe `wb.save()` calls are absorbed. That instance attribute would
    shadow the REAL openpyxl save here, producing a missing/truncated file.
    We therefore invoke the genuine bound method captured from the class.
    """
    save_path = Path(save_path)
    pre_existing = save_path.is_file()
    snapshot_bytes = save_path.read_bytes() if pre_existing else None
    # Invoke the GENUINE class-level save: an explicit class-attribute call
    # bypasses any no-op instance-level `save` shim installed by the sandbox.
    real_save = getattr(type(obj), "save", None)
    if real_save is None:
        real_save = obj.save
    try:
        real_save(obj, str(save_path))
    except Exception as save_err:
        if pre_existing:
            save_path.write_bytes(snapshot_bytes)
        else:
            save_path.unlink(missing_ok=True)
        return {
            "success": False,
            "error": (
                f"Save failed before validation could run: {save_err}\n"
                "The file was restored to its pre-save state. Fix the recipe and retry."
            ),
        }

    ok, reason = _validate_office_file(save_path)
    if ok:
        return {
            "success": True,
            "saved_path": str(save_path),
            "file_name": save_path.name,
        }

    if pre_existing:
        save_path.write_bytes(snapshot_bytes)
        restoration_note = "The file was restored to its pre-save state (pre-save in-memory snapshot)."
    else:
        save_path.unlink(missing_ok=True)
        restoration_note = "The new file was deleted (no previous version to restore)."

    ASCIIColors.error(f"[Excel Guard] {save_path.name} failed validation ({reason}) — rolled back.")
    return {
        "success": False,
        "error": (
            f"Office Document Integrity Failure [{OFFICE_GUARD_FAILURE_CODE}]: the saved workbook "
            f"failed strict validation and was rolled back.\n"
            f"File: {save_path.name}\n"
            f"Reason: {reason}\n"
            f"{restoration_note}\n"
            f"{OFFICE_GUARD_GUIDANCE}"
        ),
    }


def _exec_recipe(recipe: str, namespace: Dict[str, Any], source_name: str) -> Optional[Dict[str, Any]]:
    """Validates and executes a recipe in the sandboxed namespace."""
    if not recipe or not recipe.strip():
        return None
    if len(recipe) > _MAX_RECIPE_CHARS:
        return {
            "success": False,
            "error": f"Recipe exceeds the maximum length of {_MAX_RECIPE_CHARS} characters.",
        }
    try:
        _check_recipe(recipe)
    except ValueError as forbidden_err:
        return {"success": False, "error": str(forbidden_err)}

    dedented = textwrap.dedent(recipe)
    namespace["__builtins__"] = _build_safe_builtins()
    try:
        exec(compile(dedented, source_name, "exec"), namespace)
        return None
    except Exception:
        annotated = _annotate_traceback(traceback.format_exc(), dedented, source_name)
        if "save(" in annotated and "missing" in annotated:
            annotated += (
                "\nHint: the tool saves the file automatically after your recipe — "
                "remove the save() call."
            )
        return {"success": False, "error": f"Recipe execution error:\n{annotated}"}


def _resolve_working_dir(discussion_instance: Optional[Any]) -> Path:
    """Resolves the directory files are saved to / read from (discussion workspace or CWD)."""
    if discussion_instance is not None:
        workspace = getattr(discussion_instance, "workspace_data_path", None)
        if workspace:
            return Path(workspace)
    return Path(".")


def _sanitize_file_name(file: str, default: str) -> str:
    """Normalizes a file name, strips path components, and applies the default when empty."""
    name = (file or "").strip()
    if not name:
        return default
    name = Path(name).name
    if not name or name in (".", ".."):
        return default
    return name


def _create_root_object(suffix: str, namespace: Dict[str, Any]) -> Any:
    """Creates a fresh workbook and binds it into the namespace as wb/ws.

    Uses the namespace's shimmed Workbook factory so the pre-created root is
    save-absorbed exactly like recipe-constructed workbooks.
    """
    if suffix in (".xlsx", ".xlsm"):
        wb = namespace["Workbook"]()
        ws = wb.active
        namespace["wb"] = wb
        namespace["ws"] = ws
        return wb
    raise ValueError(f"Unsupported Excel file extension: {suffix}")


def _open_existing(path: Path, namespace: Dict[str, Any]) -> Any:
    """Opens an existing workbook by suffix and binds it into the namespace."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        wb = namespace["load_workbook"](str(path), keep_vba=(suffix == ".xlsm"))
        namespace["wb"] = wb
        namespace["ws"] = wb.active
        return wb
    raise ValueError(f"Unsupported Excel file extension: {suffix}")


def _pick_effective_root(namespace: Dict[str, Any], suffix: str, original_root: Any) -> Any:
    """Returns the recipe-rebound root alias when the recipe replaced the
    pre-bound root object (e.g. `wb = Workbook()`), else the original root."""
    aliases = {".xlsx": ("wb",), ".xlsm": ("wb",)}
    for alias in aliases.get(suffix, ()):
        candidate = namespace.get(alias)
        if candidate is not None and candidate is not original_root:
            return candidate
    return original_root


def _build_result(save_result: Dict[str, Any], save_path: Path) -> Dict[str, Any]:
    """Shapes a successful save result into the standard result contract."""
    return {
        "success": True,
        "workspace_file": str(save_path.resolve()),
        "context_files": [str(save_path.resolve())],
        "file_name": save_path.name,
    }


def _require_suffix(
    file: str,
    allowed_suffixes: Tuple[str, ...],
    tool_name: str,
) -> Optional[Dict[str, Any]]:
    """Returns a structured failure dict when the file suffix is outside the
    tool's format family, else None (proceed)."""
    actual = Path(str(file or "")).suffix.lower()
    if actual not in allowed_suffixes:
        return {
            "success": False,
            "error": (
                f"{tool_name} requires a {'/'.join(allowed_suffixes)} file, "
                f"got '{actual or '(no extension)'}'. Use the tool matching "
                f"this file's format instead."
            ),
            "code": "wrong_file_type_for_tool",
        }
    return None


def _read_xlsx(path: Path) -> str:
    """Builds a markdown inventory of a workbook (sheets, dimensions, cell previews, formulas)."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=False)
    lines = [f"# Workbook: {path.name}", f"Sheets: {len(wb.sheetnames)}", ""]
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"## Sheet: {sheet_name} ({ws.max_row} rows x {ws.max_column} cols)")
        max_preview_rows = min(ws.max_row, 20)
        max_preview_cols = min(ws.max_column, 10)
        formula_cells = []
        for row in ws.iter_rows(min_row=1, max_row=max_preview_rows, max_col=max_preview_cols):
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append(f"{cell.coordinate}: {cell.value}")
        if formula_cells:
            lines.append("Formulas (first rows):")
            lines.extend(f"- {entry}" for entry in formula_cells[:20])
        lines.append("Preview (values, first 20 rows x 10 cols):")
        lines.append("```")
        for row in ws.iter_rows(min_row=1, max_row=max_preview_rows, max_col=max_preview_cols, values_only=True):
            lines.append(" | ".join("" if v is None else str(v) for v in row))
        lines.append("```")
        lines.append("")
    return "\n".join(lines)


def tool_build_excel(
    file: str = "workbook.xlsx",
    recipe: str = "",
    discussion_instance: Optional[Any] = None,
    lollms_client_instance: Optional[Any] = None,
) -> Dict[str, Any]:
    """Creates a new Excel workbook (.xlsx/.xlsm) from a sandboxed openpyxl recipe.

    === EXCEL DOCTRINE (CRITICAL — READ BEFORE WRITING THE RECIPE) ===
    NEVER write computed results manually. ALWAYS write FORMULAS into cells so
    Excel computes them:
        ws["F2"] = "=SUM(B2:E2)"
        ws["G2"] = "=AVERAGE(B2:E2)"
        ws["H2"] = '=IF(F2>100,"High","Low")'
        ws["I2"] = "=VLOOKUP(A2,$K$2:$L$10,2,FALSE)"
    Use SUM, AVERAGE, MIN, MAX, COUNT, IF, VLOOKUP, INDEX/MATCH, CONCATENATE,
    etc. Literal values are ONLY for input data (raw numbers, dates, labels).
    Writing a computed number by hand (e.g. ws["F2"] = 42 when F2 should be
    SUM(B2:E2)) is a CRITICAL ERROR.

    Chart caveat: when a chart's Reference points at formula cells, openpyxl
    writes formulas WITHOUT cached values — some desktop viewers show an empty
    chart until Excel recalculates. Formulas remain the preferred approach;
    precomputed literals are acceptable only when the user explicitly needs
    chart rendering without recalculation.

    The root objects are pre-bound as `wb` (workbook) and `ws` (active sheet).
    Available symbols: Workbook, load_workbook, PieChart, LineChart, BarChart,
    Reference, Font, PatternFill, Alignment, Border, Side, get_column_letter,
    CellIsRule, ColorScaleRule, DataValidation, media(), list_media().
    Recipes are plain Python: imports are allowed ONLY from the whitelist
    (openpyxl.* and safe stdlib modules); all core design symbols are also
    pre-imported. The tool saves automatically; calling wb.save() inside the
    recipe is a no-op — remove it. Also available: tables, conditional
    formatting (CellIsRule/ColorScaleRule), data validation (DataValidation),
    freeze panes (ws.freeze_panes), column widths (ws.column_dimensions),
    number formats (cell.number_format).

    Args:
        file: Target .xlsx/.xlsm file name (default "workbook.xlsx").
        recipe: Plain Python recipe code mutating the pre-bound wb/ws objects.
        discussion_instance: Active discussion session instance.
        lollms_client_instance: Active client instance.
    """
    suffix_failure = _require_suffix(file, _XLSX_SUFFIXES, "tool_build_excel")
    if suffix_failure is not None:
        return suffix_failure
    try:
        file_name = _sanitize_file_name(file, "workbook.xlsx")
        suffix = Path(file_name).suffix.lower()
        if suffix not in GUARDED_EXTENSIONS:
            return {
                "success": False,
                "error": f"Unsupported Excel file extension '{suffix}'. Supported: {', '.join(GUARDED_EXTENSIONS)}.",
            }

        working_dir = _resolve_working_dir(discussion_instance)
        namespace = _build_recipe_namespace(working_dir=working_dir)
        try:
            root_obj = _create_root_object(suffix, namespace)
        except KeyError as missing_symbol:
            return {
                "success": False,
                "error": (
                    f"openpyxl is not installed (missing symbol: {missing_symbol}). "
                    "Ask the user to install it and retry."
                ),
            }

        try:
            root_obj.save = lambda *args, **kwargs: None  # the tool performs the real validated save
        except Exception:
            pass

        exec_failure = _exec_recipe(recipe, namespace, "<excel_recipe>")
        if exec_failure is not None:
            return exec_failure

        root_obj = _pick_effective_root(namespace, suffix, root_obj)

        save_path = working_dir / file_name
        save_result = _save_and_verify(root_obj, save_path)
        if not save_result["success"]:
            return save_result
        return _build_result(save_result, save_path)
    except Exception as unexpected_err:
        return {
            "success": False,
            "error": f"Unexpected failure building Excel workbook: {unexpected_err}",
        }


def tool_edit_excel(
    file: str,
    recipe: str = "",
    discussion_instance: Optional[Any] = None,
    lollms_client_instance: Optional[Any] = None,
) -> Dict[str, Any]:
    """Edits an existing Excel workbook (.xlsx/.xlsm) in place via a sandboxed openpyxl recipe.

    === EXCEL DOCTRINE (CRITICAL — READ BEFORE WRITING THE RECIPE) ===
    NEVER write computed results manually. ALWAYS write FORMULAS into cells so
    Excel computes them:
        ws["F2"] = "=SUM(B2:E2)"
        ws["G2"] = "=AVERAGE(B2:E2)"
        ws["H2"] = '=IF(F2>100,"High","Low")'
    Use SUM, AVERAGE, MIN, MAX, COUNT, IF, VLOOKUP, INDEX/MATCH, CONCATENATE,
    etc. Literal values are ONLY for input data. Writing a computed number by
    hand instead of a formula is a CRITICAL ERROR.

    The opened workbook is pre-bound as `wb` and its active sheet as `ws`
    (.xlsm files are opened with keep_vba=True). Available symbols: Workbook,
    load_workbook, PieChart, LineChart, BarChart, Reference, Font, PatternFill,
    Alignment, Border, Side, get_column_letter, CellIsRule, ColorScaleRule,
    DataValidation, media(), list_media(). Recipes are plain Python: imports
    are allowed ONLY from the whitelist (openpyxl.* and safe stdlib modules);
    all core design symbols are also pre-imported. The tool saves
    automatically; calling wb.save() inside the recipe is a no-op — remove it.
    The save is transactional: on validation failure the pre-edit bytes are
    restored.

    Args:
        file: Existing .xlsx/.xlsm file to edit (relative to workspace or absolute path).
        recipe: Plain Python recipe code mutating the opened wb/ws objects.
        discussion_instance: Active discussion session instance.
        lollms_client_instance: Active client instance.
    """
    suffix_failure = _require_suffix(file, _XLSX_SUFFIXES, "tool_edit_excel")
    if suffix_failure is not None:
        return suffix_failure
    try:
        if not file or not str(file).strip():
            return {"success": False, "error": "No file provided for editing."}

        working_dir = _resolve_working_dir(discussion_instance)
        candidate = Path(str(file).strip())
        file_path = candidate if candidate.is_absolute() else working_dir / candidate
        if not file_path.is_file():
            return {"success": False, "error": f"File not found: {file_path}"}

        suffix = file_path.suffix.lower()
        if suffix not in GUARDED_EXTENSIONS:
            return {
                "success": False,
                "error": f"Unsupported Excel file extension '{suffix}'. Supported: {', '.join(GUARDED_EXTENSIONS)}.",
            }

        namespace = _build_recipe_namespace(working_dir=working_dir)
        try:
            root_obj = _open_existing(file_path, namespace)
        except KeyError as missing_symbol:
            return {
                "success": False,
                "error": (
                    f"openpyxl is not installed (missing symbol: {missing_symbol}). "
                    "Ask the user to install it and retry."
                ),
            }

        try:
            root_obj.save = lambda *args, **kwargs: None  # the tool performs the real validated save
        except Exception:
            pass

        exec_failure = _exec_recipe(recipe, namespace, "<excel_recipe>")
        if exec_failure is not None:
            return exec_failure

        root_obj = _pick_effective_root(namespace, suffix, root_obj)

        save_result = _save_and_verify(root_obj, file_path)
        if not save_result["success"]:
            return save_result
        return _build_result(save_result, file_path)
    except Exception as unexpected_err:
        return {
            "success": False,
            "error": f"Unexpected failure editing Excel workbook: {unexpected_err}",
        }


_MD_SEPARATOR_CELL_RE = re.compile(r"^\s*\|?\s*:?-{1,}:?\s*$")
_MD_MAX_COLUMN_WIDTH = 50


def _split_markdown_table_row(line: str) -> List[str]:
    """Splits a markdown pipe-table row into stripped cell values."""
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def _is_markdown_separator_row(line: str) -> bool:
    """Detects a pipe-table separator row (|---|:---:|---|)."""
    if "|" not in line:
        return False
    cells = _split_markdown_table_row(line)
    if not cells:
        return False
    return all(_MD_SEPARATOR_CELL_RE.match(cell) for cell in cells if cell != "")


def _coerce_markdown_cell(value: str) -> Any:
    """Types a cell value: clean int/float strings become numbers, all else strings.

    No invented data: '-', '98.7%', '-2π à 2π' and empty strings stay strings
    exactly as written; only unambiguous numeric literals are converted.
    """
    text = (value or "").strip()
    if not text:
        return ""
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        pass
    return text


def _parse_markdown_tables(content: str) -> List[Dict[str, Any]]:
    """Parses markdown into ordered sections: headings, pipe tables, plain text.

    Returns a list of blocks in document order, each shaped as
    {"title": str|None, "headers": List[str], "rows": List[List], "text_lines": List[str]}.
    Code fences are ignored; a pipe row above a separator row becomes the
    header row of a table; remaining pipe rows are data rows.
    """
    blocks: List[Dict[str, Any]] = []
    current: Dict[str, Any] = {"title": None, "headers": [], "rows": [], "text_lines": []}

    def _flush() -> None:
        nonlocal current
        if current["title"] is not None or current["headers"] or current["rows"] or current["text_lines"]:
            blocks.append(current)
        current = {"title": None, "headers": [], "rows": [], "text_lines": []}

    inside_code_fence = False
    pending_pipe_rows: List[List[str]] = []

    def _absorb_pending() -> None:
        nonlocal pending_pipe_rows
        for cells in pending_pipe_rows:
            current["text_lines"].append(" | ".join(cells))
        pending_pipe_rows = []

    for raw_line in (content or "").splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            inside_code_fence = not inside_code_fence
            continue
        if inside_code_fence:
            continue
        if not stripped:
            _absorb_pending()
            continue
        if stripped.startswith("#"):
            _absorb_pending()
            _flush()
            current["title"] = stripped.lstrip("#").strip()
            continue
        if _is_markdown_separator_row(stripped):
            if pending_pipe_rows:
                current["headers"] = pending_pipe_rows.pop(0)
                current["rows"].extend(pending_pipe_rows)
                pending_pipe_rows = []
            continue
        if stripped.startswith("|") and stripped.count("|") >= 1:
            pending_pipe_rows.append(_split_markdown_table_row(stripped))
            continue
        _absorb_pending()
        current["text_lines"].append(stripped)

    _absorb_pending()
    _flush()
    return blocks


def tool_excel_from_markdown(
    file: str = "generated_report.xlsx",
    content: str = "",
    discussion_instance: Optional[Any] = None,
    lollms_client_instance: Optional[Any] = None,
) -> Dict[str, Any]:
    """Creates a real Excel (.xlsx) workbook from a markdown document (headings + pipe tables).

    Excel files are BINARY OOXML workbooks (ZIP of XML parts) — this tool
    converts markdown into genuine openpyxl worksheets with styled section
    titles, bold header rows and typed cells (clean numbers become numbers,
    everything else stays text); it NEVER writes raw text or markdown into an
    .xlsx file. Use it whenever the source material is markdown or structured
    text instead of a Python recipe.

    Args:
        file: Target .xlsx/.xlsm file name (default "generated_report.xlsx").
        content: Markdown payload: '#'/'##' headings become section title rows,
            pipe tables become styled header + data rows, plain lines become
            single-column text rows.
        discussion_instance: Active discussion session instance.
        lollms_client_instance: Active client instance.
    """
    suffix_failure = _require_suffix(file, _XLSX_SUFFIXES, "tool_excel_from_markdown")
    if suffix_failure is not None:
        return suffix_failure
    try:
        if not content or not str(content).strip():
            return {
                "success": False,
                "error": "No markdown content provided. Pass the document text (headings and/or pipe tables) via the 'content' parameter.",
            }

        file_name = _sanitize_file_name(file, "generated_report.xlsx")
        suffix = Path(file_name).suffix.lower()
        if suffix not in GUARDED_EXTENSIONS:
            return {
                "success": False,
                "error": f"Unsupported Excel file extension '{suffix}'. Supported: {', '.join(GUARDED_EXTENSIONS)}.",
            }

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        blocks = _parse_markdown_tables(str(content))
        if not blocks:
            return {
                "success": False,
                "error": "The provided content contains no usable markdown (no headings, tables or text lines).",
            }

        title_font = Font(bold=True, size=13)
        header_font = Font(bold=True)
        header_fill = PatternFill(
            start_color="FFD9E1F2", end_color="FFD9E1F2", fill_type="solid"
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        column_widths: Dict[int, int] = {}

        def _track(column_index: int, value: Any) -> None:
            length = len(str(value)) if value is not None else 0
            if length > column_widths.get(column_index, 0):
                column_widths[column_index] = length

        for block in blocks:
            if block["title"] is not None:
                ws.cell(row=ws.max_row + 1 if ws.max_row > 1 or ws["A1"].value is not None else 1, column=1, value=block["title"]).font = title_font
                _track(1, block["title"])
            for text_line in block["text_lines"]:
                ws.cell(row=ws.max_row + 1, column=1, value=text_line)
                _track(1, text_line)
            if block["headers"] or block["rows"]:
                header_row_index = ws.max_row + 1 if ws["A1"].value is not None or ws.max_row > 1 else 1
                for column_index, header in enumerate(block["headers"], start=1):
                    cell = ws.cell(row=header_row_index, column=column_index, value=header)
                    cell.font = header_font
                    cell.fill = header_fill
                    _track(column_index, header)
                for data_cells in block["rows"]:
                    row_index = ws.max_row + 1
                    for column_index, raw_value in enumerate(data_cells, start=1):
                        typed_value = _coerce_markdown_cell(raw_value)
                        ws.cell(row=row_index, column=column_index, value=typed_value)
                        _track(column_index, raw_value)
                ws.cell(row=ws.max_row + 1, column=1, value=None)  # spacer row between sections

        for column_index, width in column_widths.items():
            ws.column_dimensions[get_column_letter(column_index)].width = min(width + 2, _MD_MAX_COLUMN_WIDTH)

        working_dir = _resolve_working_dir(discussion_instance)
        save_path = working_dir / file_name
        save_result = _save_and_verify(wb, save_path)
        if not save_result["success"]:
            return save_result
        return _build_result(save_result, save_path)
    except Exception as unexpected_err:
        return {
            "success": False,
            "error": f"Unexpected failure converting markdown to Excel workbook: {unexpected_err}",
        }


def tool_read_excel(
    file: str,
    discussion_instance: Optional[Any] = None,
    lollms_client_instance: Optional[Any] = None,
) -> Dict[str, Any]:
    """Reads an existing Excel workbook (.xlsx/.xlsm) and returns a markdown inventory.

    Returns sheets, dimensions, cell previews and the formulas present in the
    first rows. Validates the file first; corrupt files return a structured
    failure instead of a crash. No recipe is executed and nothing is saved.
    """
    try:
        if not file or not str(file).strip():
            return {"success": False, "error": "No file provided for reading."}

        working_dir = _resolve_working_dir(discussion_instance)
        candidate = Path(str(file).strip())
        file_path = candidate if candidate.is_absolute() else working_dir / candidate
        if not file_path.is_file():
            return {"success": False, "error": f"File not found: {file_path}"}

        suffix = file_path.suffix.lower()
        if suffix not in _XLSX_SUFFIXES:
            return {
                "success": False,
                "error": f"tool_read_excel requires a {'/'.join(_XLSX_SUFFIXES)} file, got '{suffix or '(no extension)'}'.",
            }

        ok, reason = _validate_office_file(file_path)
        if not ok:
            return {
                "success": False,
                "error": (
                    f"Office Document Integrity Failure [{OFFICE_GUARD_FAILURE_CODE}]: "
                    f"'{file_path.name}' is corrupt and cannot be read safely.\n"
                    f"Reason: {reason}"
                ),
            }

        inventory = _read_xlsx(file_path)
        return {
            "success": True,
            "file_name": file_path.name,
            "workspace_file": str(file_path.resolve()),
            "inventory": inventory,
        }
    except Exception as unexpected_err:
        return {
            "success": False,
            "error": f"Unexpected failure reading Excel workbook: {unexpected_err}",
        }


def get_tools() -> List[Dict[str, Any]]:
    """Returns the tool specifications for the Excel Tools toolset."""
    return [
        {
            "name": "tool_build_excel",
            "description": (
                "Creates a new Excel workbook (.xlsx/.xlsm) from a sandboxed openpyxl recipe. "
                "DOCTRINE: NEVER write computed results manually — ALWAYS write formulas so Excel "
                "computes them: ws['F2']='=SUM(B2:E2)', ws['G2']='=AVERAGE(B2:E2)', "
                "ws['H2']='=IF(F2>100,\"High\",\"Low\")'. Literals only for input data. Root objects "
                "pre-bound as wb/ws. Symbols: Workbook, load_workbook, PieChart, LineChart, BarChart, "
                "Reference, Font, PatternFill, Alignment, Border, Side, get_column_letter, CellIsRule, "
                "ColorScaleRule, DataValidation, media(), list_media(). Imports allowed ONLY from the "
                "whitelist (openpyxl.*, safe stdlib). Do NOT call save() — the tool saves and validates "
                "automatically. Chart note: charts referencing formula cells may appear empty in some "
                "viewers until Excel recalculates."
            ),
            "parameters": [
                {
                    "name": "file",
                    "type": "str",
                    "description": "Target .xlsx/.xlsm file name (default 'workbook.xlsx').",
                    "optional": True,
                },
                {
                    "name": "recipe",
                    "type": "str",
                    "description": "Plain Python recipe mutating the pre-bound wb/ws. FORMULAS ONLY for computed cells (=SUM, =AVERAGE, =IF, =VLOOKUP...). Imports allowed only from the whitelist (openpyxl.*, safe stdlib). Do NOT call save().",
                    "optional": True,
                },
            ],
            "output": [
                {"name": "success", "type": "bool", "description": "Whether the operation succeeded."},
                {"name": "workspace_file", "type": "str", "description": "Absolute path to the saved file."},
                {"name": "context_files", "type": "list", "description": "List of context files (the saved file)."},
                {"name": "file_name", "type": "str", "description": "Name of the saved file."},
                {"name": "error", "type": "str", "description": "Error message if failed."},
            ],
            "callable": tool_build_excel,
            "_source": "excel_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_edit_excel",
            "description": (
                "Edits an existing Excel workbook (.xlsx/.xlsm) in place via a sandboxed openpyxl "
                "recipe. DOCTRINE: NEVER write computed results manually — ALWAYS write formulas so "
                "Excel computes them (=SUM, =AVERAGE, =IF, =VLOOKUP...). Literals only for input data. "
                "Opened workbook pre-bound as wb, active sheet as ws (.xlsm opened with keep_vba=True). "
                "Imports allowed ONLY from the whitelist (openpyxl.*, safe stdlib). Do NOT call save() — "
                "the tool saves and validates automatically. Transactional save: pre-edit bytes restored "
                "on validation failure."
            ),
            "parameters": [
                {
                    "name": "file",
                    "type": "str",
                    "description": "Existing .xlsx/.xlsm file to edit (relative to workspace or absolute path).",
                },
                {
                    "name": "recipe",
                    "type": "str",
                    "description": "Plain Python recipe mutating the opened wb/ws. FORMULAS ONLY for computed cells. Imports allowed only from the whitelist (openpyxl.*, safe stdlib). Do NOT call save().",
                    "optional": True,
                },
            ],
            "output": [
                {"name": "success", "type": "bool", "description": "Whether the operation succeeded."},
                {"name": "workspace_file", "type": "str", "description": "Absolute path to the saved file."},
                {"name": "context_files", "type": "list", "description": "List of context files (the saved file)."},
                {"name": "file_name", "type": "str", "description": "Name of the saved file."},
                {"name": "error", "type": "str", "description": "Error message if failed."},
            ],
            "callable": tool_edit_excel,
            "_source": "excel_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_read_excel",
            "description": (
                "Reads an existing Excel workbook (.xlsx/.xlsm) and returns a markdown inventory "
                "(sheets, dimensions, cell previews, formulas). Validates the file first; corrupt "
                "files return a structured failure."
            ),
            "parameters": [
                {
                    "name": "file",
                    "type": "str",
                    "description": "Existing .xlsx/.xlsm file to read (relative to workspace or absolute path).",
                },
            ],
            "output": [
                {"name": "success", "type": "bool", "description": "Whether the operation succeeded."},
                {"name": "inventory", "type": "str", "description": "Markdown inventory of the workbook contents."},
                {"name": "file_name", "type": "str", "description": "Name of the file read."},
                {"name": "workspace_file", "type": "str", "description": "Absolute path to the file."},
                {"name": "error", "type": "str", "description": "Error message if failed."},
            ],
            "callable": tool_read_excel,
            "_source": "excel_tools",
            "_source_file": __file__,
        },
        {
            "name": "tool_excel_from_markdown",
            "description": (
                "Creates a real Excel (.xlsx) workbook from a markdown document (headings + pipe "
                "tables). Excel files are BINARY OOXML workbooks — this tool converts markdown into "
                "genuine openpyxl worksheets with styled section titles, bold header rows and typed "
                "cells (clean numbers become numbers, everything else stays text); it NEVER writes "
                "raw text or markdown into an .xlsx file. Use it whenever the source material is "
                "markdown or structured text instead of a Python recipe."
            ),
            "parameters": [
                {
                    "name": "file",
                    "type": "str",
                    "description": "Target .xlsx/.xlsm file name (default 'generated_report.xlsx').",
                    "optional": True,
                },
                {
                    "name": "content",
                    "type": "str",
                    "description": "Markdown payload: '#'/'##' headings become section title rows, pipe tables become styled header + data rows, plain lines become single-column text rows.",
                },
            ],
            "output": [
                {"name": "success", "type": "bool", "description": "Whether the operation succeeded."},
                {"name": "workspace_file", "type": "str", "description": "Absolute path to the saved file."},
                {"name": "context_files", "type": "list", "description": "List of context files (the saved file)."},
                {"name": "file_name", "type": "str", "description": "Name of the saved file."},
                {"name": "error", "type": "str", "description": "Error message if failed."},
            ],
            "callable": tool_excel_from_markdown,
            "_source": "excel_tools",
            "_source_file": __file__,
        },
    ]