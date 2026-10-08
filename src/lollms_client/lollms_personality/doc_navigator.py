"""
doc_navigator.py
Bibliographic documentation system for Handbags.

Doctrine (do not deviate):
- Handbag knowledge uses hierarchical INDEX navigation + fuzzy grep, NEVER vector embeddings.
- One document = one directory; markdown heading depth defines the tree depth.
- A chunk is one complete heading section (title + full body); never split mid-content.
- Every loaded/searched fragment carries its full canonical path (provenance-native).
- Over-budget loads degrade to the target's index + per-child estimates, never a hard failure.
- Metadata extraction is deterministic-first; an LLM is used only behind an explicit use_llm flag.
"""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple, Union

from ascii_colors import ASCIIColors


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_TXT_CHAPTER_RE = re.compile(r"(?i)^(chapter|part|section)\s+[\dIVXLC]+", re.IGNORECASE)
_ABSTRACT_RE = re.compile(
    r"(?im)^#{0,6}\s*(?:abstract|summary|résumé)\s*#*\s*\n+(.*?)(?=\n#{1,6}\s|\Z)",
    re.DOTALL,
)
_YEAR_RE = re.compile(r"\b(19\d{2}|20\d{2})\b")
_AUTHOR_LINE_RE = re.compile(r"(?im)^(?:by|authors?|written\s+by|édité\s+par)\s*[:\-]?\s*(.{3,150})$")
_NAMES_LINE_RE = re.compile(r"^[A-Z][A-Za-z.\-' ]+(?:\s*(?:,|&|and|et)\s*[A-Z][A-Za-z.\-' ]+){1,5}$")

_WINDOWS_RESERVED = {
    "CON", "PRN", "AUX", "NUL",
    *{f"COM{i}" for i in range(1, 10)},
    *{f"LPT{i}" for i in range(1, 10)},
}

_MAX_SLUG_LEN = 80
_MAX_GLIMPSE_DOCS = 40
_PEEK_WINDOW_LINES = 80
_DEFAULT_MAX_CHARS = 24000


def _estimate_tokens(text: str) -> int:
    if not text:
        return 0
    return max(1, len(text) // 4)


def _safe_slug(title: str, max_len: int = _MAX_SLUG_LEN) -> str:
    normalized = unicodedata.normalize("NFKD", title or "")
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-zA-Z0-9_-]+", "_", ascii_text).strip("_").lower()
    if not slug:
        slug = "untitled"
    if len(slug) > max_len:
        truncated = slug[:max_len].rsplit("_", 1)[0] or slug[:max_len]
        digest = hashlib.md5((title or "").encode("utf-8", "ignore")).hexdigest()[:6]
        slug = f"{truncated}_{digest}"
    if slug.upper() in _WINDOWS_RESERVED:
        slug = f"_{slug}"
    return slug


def _unique_dir(parent: Path, slug: str) -> Path:
    candidate = parent / slug
    if not candidate.exists():
        return candidate
    index = 1
    while True:
        candidate = parent / f"{slug}_{index}"
        if not candidate.exists():
            return candidate
        index += 1


def _split_frontmatter(text: str) -> Tuple[str, Dict[str, str]]:
    if text.lstrip().startswith("---"):
        parts = text.split("---", 2)
        if len(parts) >= 3:
            raw_fm = parts[1]
            body = parts[2].lstrip("\r\n")
            fm: Dict[str, str] = {}
            for line in raw_fm.splitlines():
                if ":" in line:
                    key, _, value = line.partition(":")
                    key = key.strip().lower()
                    value = value.strip().strip("'\"")
                    if key and value:
                        fm[key] = value
            return body, fm
    return text, {}


def _yaml_quote(value: str) -> str:
    cleaned = (value or "").replace("\n", " ").replace('"', "'").strip()
    return f'"{cleaned}"' if cleaned else '""'


def _frontmatter_block(fields: Dict[str, Any]) -> str:
    lines = ["---"]
    for key, value in fields.items():
        if value is None:
            continue
        if isinstance(value, str):
            lines.append(f"{key}: {_yaml_quote(value)}")
        else:
            lines.append(f"{key}: {value}")
    lines.append("---")
    return "\n".join(lines)


# ───────────────────────────────────────────────────────────────────────────
# Section model & markdown parsing
# ───────────────────────────────────────────────────────────────────────────

@dataclass
class _Section:
    title: str
    level: int
    content: str
    children: List["_Section"] = field(default_factory=list)


def _parse_markdown_sections(md_text: str) -> Tuple[List[Tuple[int, str, str]], str]:
    flat: List[Tuple[int, str, str]] = []
    preamble_lines: List[str] = []
    current_level: Optional[int] = None
    current_title: Optional[str] = None
    buffer: List[str] = []

    def _flush() -> None:
        nonlocal current_level, current_title, buffer
        if current_title is not None:
            flat.append((current_level or 1, current_title, "\n".join(buffer).strip()))
        current_level = None
        current_title = None
        buffer = []

    for line in md_text.splitlines():
        match = _HEADING_RE.match(line)
        if match:
            _flush()
            current_level = len(match.group(1))
            current_title = match.group(2).strip()
        elif current_title is not None:
            buffer.append(line)
        else:
            preamble_lines.append(line)
    _flush()

    return flat, "\n".join(preamble_lines).strip()


def _build_section_tree(flat: List[Tuple[int, str, str]]) -> List[_Section]:
    if not flat:
        return []
    unique_levels = sorted({level for level, _, _ in flat})
    level_map = {level: idx + 1 for idx, level in enumerate(unique_levels)}
    roots: List[_Section] = []
    stack: List[Tuple[int, _Section]] = []
    for level, title, content in flat:
        depth = level_map[level]
        node = _Section(title=title, level=depth, content=content)
        while stack and stack[-1][0] >= depth:
            stack.pop()
        if not stack:
            roots.append(node)
        else:
            stack[-1][1].children.append(node)
        stack.append((depth, node))
    return roots


def _subtree_tokens(section: _Section) -> int:
    total = _estimate_tokens(section.content)
    for child in section.children:
        total += _subtree_tokens(child)
    return total


def _emit_section_tree(sections: List[_Section], parent_dir: Path) -> List[Dict[str, Any]]:
    entries: List[Dict[str, Any]] = []
    for index, section in enumerate(sections, 1):
        prefix = f"{index:02d}_"
        slug = _safe_slug(section.title)
        if section.children:
            section_dir = parent_dir / f"{prefix}{slug}"
            section_dir.mkdir(parents=True, exist_ok=True)
            child_entries = _emit_section_tree(section.children, section_dir)
            _write_sub_index(section, section_dir, child_entries)
            entries.append({
                "name": section_dir.name,
                "title": section.title,
                "tokens": _subtree_tokens(section),
                "is_dir": True,
                "children": len(child_entries),
            })
        else:
            body = f"# {section.title}\n\n{section.content}\n".strip() + "\n"
            file_path = parent_dir / f"{prefix}{slug}.md"
            file_path.write_text(body, encoding="utf-8")
            entries.append({
                "name": file_path.name,
                "title": section.title,
                "tokens": _estimate_tokens(body),
                "is_dir": False,
                "children": 0,
            })
    return entries


def _write_sub_index(section: _Section, section_dir: Path, child_entries: List[Dict[str, Any]]) -> None:
    lines = [f"# {section.title}", ""]
    if section.content.strip():
        lines.extend([section.content.strip(), ""])
    lines.append(f"## Subsections of {section.title}")
    for entry in child_entries:
        if entry["is_dir"]:
            lines.append(f"- [{entry['title']}]({entry['name']}/) (~{entry['tokens']} tokens, {entry['children']} subsection(s))")
        else:
            lines.append(f"- [{entry['title']}]({entry['name']}) (~{entry['tokens']} tokens)")
    lines.append("")
    (section_dir / "INDEX.md").write_text("\n".join(lines), encoding="utf-8")


# ───────────────────────────────────────────────────────────────────────────
# Format converters (deterministic, lazy optional imports)
# ───────────────────────────────────────────────────────────────────────────

def _convert_txt(raw: str) -> Tuple[str, Dict[str, str]]:
    out_lines: List[str] = []
    for line in raw.splitlines():
        stripped = line.strip()
        is_chapter = bool(_TXT_CHAPTER_RE.match(stripped))
        is_caps_heading = (
            3 < len(stripped) <= 80
            and stripped == stripped.upper()
            and any(c.isalpha() for c in stripped)
            and not stripped.endswith((".", ",", ";", ":"))
        )
        if is_chapter:
            out_lines.append(f"# {stripped}")
        elif is_caps_heading:
            out_lines.append(f"## {stripped}")
        else:
            out_lines.append(line)
    return "\n".join(out_lines), {}


def _convert_pdf(file_path: Path) -> Tuple[str, Dict[str, str]]:
    try:
        import fitz
    except ImportError as import_err:
        raise ImportError(
            "PDF ingestion requires PyMuPDF. Install it with: pip install pymupdf"
        ) from import_err

    doc = fitz.open(str(file_path))
    try:
        metadata = doc.metadata or {}
        hints = {
            "title": (metadata.get("title") or "").strip(),
            "authors": (metadata.get("author") or "").strip(),
        }
        size_counter: Counter = Counter()
        lines: List[Tuple[float, str]] = []
        for page in doc:
            page_dict = page.get_text("dict")
            for block in page_dict.get("blocks", []):
                for line_info in block.get("lines", []):
                    spans = line_info.get("spans", [])
                    text = "".join(span.get("text", "") for span in spans).strip()
                    if not text:
                        continue
                    size = max((span.get("size", 10.0) for span in spans), default=10.0)
                    rounded = round(size, 1)
                    size_counter[rounded] += 1
                    lines.append((rounded, text))

        body_size = size_counter.most_common(1)[0][0] if size_counter else 10.0
        heading_sizes = sorted(
            {size for size in size_counter if size >= body_size * 1.18},
            reverse=True,
        )[:3]
        level_of = {size: idx + 1 for idx, size in enumerate(heading_sizes)}

        out_lines: List[str] = []
        for size, text in lines:
            if size in level_of:
                out_lines.append("")
                out_lines.append(f"{'#' * level_of[size]} {text}")
                out_lines.append("")
            else:
                out_lines.append(text)
        return "\n".join(out_lines), hints
    finally:
        doc.close()


def _convert_docx(file_path: Path) -> Tuple[str, Dict[str, str]]:
    try:
        from docx import Document
    except ImportError as import_err:
        raise ImportError(
            "DOCX ingestion requires python-docx. Install it with: pip install python-docx"
        ) from import_err

    document = Document(str(file_path))
    core = document.core_properties
    hints = {
        "title": (core.title or "").strip(),
        "authors": (core.author or "").strip(),
    }
    out_lines: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
        heading_match = re.match(r"Heading (\d)", style_name)
        if heading_match:
            out_lines.append(f"{'#' * int(heading_match.group(1))} {text}")
        elif style_name == "Title":
            out_lines.append(f"# {text}")
        else:
            out_lines.append(text)
    return "\n\n".join(out_lines), hints


def _convert_pptx(file_path: Path) -> Tuple[str, Dict[str, str]]:
    try:
        from pptx import Presentation
    except ImportError as import_err:
        raise ImportError(
            "PPTX ingestion requires python-pptx. Install it with: pip install python-pptx"
        ) from import_err

    presentation = Presentation(str(file_path))
    out_lines: List[str] = []
    for slide_index, slide in enumerate(presentation.slides, 1):
        title = ""
        title_shape = slide.shapes.title
        if title_shape is not None:
            title = (title_shape.text or "").strip()
        heading = f"Slide {slide_index}: {title}" if title else f"Slide {slide_index}"
        out_lines.append(f"## {heading}")
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if title_shape is not None and shape == title_shape:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text:
                    out_lines.append(f"- {text}")
        out_lines.append("")
    return "\n\n".join(out_lines), {}


# ───────────────────────────────────────────────────────────────────────────
# Deterministic metadata extraction
# ───────────────────────────────────────────────────────────────────────────

def _extract_title(md_text: str, hints: Dict[str, str], filename: str) -> str:
    hinted = (hints.get("title") or "").strip()
    if hinted:
        return hinted[:300]
    first_nonempty = ""
    for line in md_text.splitlines():
        if line.strip():
            first_nonempty = line.strip()
            break
    if not first_nonempty:
        return filename
    if first_nonempty.startswith("#"):
        cleaned = first_nonempty.lstrip("#").strip()
        return cleaned[:300] or filename
    cleaned = first_nonempty.strip()
    looks_like_title = (
        len(cleaned) <= 80
        and not cleaned.endswith((".", ",", "!", "?"))
        and any(c.isalpha() for c in cleaned)
    )
    if looks_like_title:
        return cleaned[:300]
    heading_match = re.search(r"^#{1,6}\s+(.+)$", md_text, re.MULTILINE)
    if heading_match:
        return heading_match.group(1).strip()[:300]
    return cleaned[:300]


def _extract_authors(md_text: str, hints: Dict[str, str]) -> str:
    hinted = (hints.get("authors") or "").strip()
    if hinted and hinted.lower() not in ("unknown", "unknown author", ""):
        return hinted[:300]
    head = md_text[:4000]
    direct = _AUTHOR_LINE_RE.search(head)
    if direct:
        candidate = direct.group(1).strip().rstrip(".,;")
        if candidate:
            return candidate[:300]
    lines = [l.strip() for l in head.splitlines() if l.strip()]
    title_idx = 0
    for idx, line in enumerate(lines[:10]):
        if line.startswith("#"):
            title_idx = idx
            break
    for line in lines[title_idx + 1: title_idx + 4]:
        if line.startswith("#"):
            break
        if _NAMES_LINE_RE.match(line) and not line.endswith("."):
            return line[:300]
    return "Unknown"


def _extract_abstract(md_text: str) -> str:
    match = _ABSTRACT_RE.search(md_text)
    if match:
        abstract = match.group(1).strip()
        if abstract and "(no abstract available)" not in abstract.lower():
            return abstract[:1200]
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", md_text) if p.strip()]
    for paragraph in paragraphs:
        if len(paragraph) > 200:
            return paragraph[:1200]
    return ""


def _extract_year(md_text: str) -> str:
    match = _YEAR_RE.search(md_text[:1000])
    if match:
        return match.group(1)
    return ""


def _llm_extract_metadata(lollms_client: Any, md_text: str, filename: str) -> Dict[str, str]:
    result: Dict[str, str] = {"title": "", "authors": "", "abstract": ""}
    if lollms_client is None:
        return result
    excerpt = md_text[:6000]
    prompt = (
        "Analyze the following document excerpt and extract its bibliographic metadata.\n"
        f"Document filename: {filename}\n\n"
        f"=== DOCUMENT EXCERPT ===\n{excerpt}\n=== END EXCERPT ===\n\n"
        "Return a JSON object with keys: title, authors, abstract.\n"
        "authors is a comma-separated string of author names, or Unknown.\n"
        "abstract is a 2-4 sentence summary of the document."
    )
    try:
        if hasattr(lollms_client, "generate_structured_content"):
            schema = {
                "title": "The document title",
                "authors": "Comma-separated author names, or Unknown",
                "abstract": "A 2-4 sentence abstract of the document",
            }
            parsed = lollms_client.generate_structured_content(
                prompt=prompt, schema=schema, temperature=0.1
            )
            if isinstance(parsed, dict):
                for key in result:
                    value = parsed.get(key)
                    if isinstance(value, str) and value.strip():
                        result[key] = value.strip()
                return result
    except Exception as ex:
        ASCIIColors.warning(f"[DocIngestor] Structured LLM metadata extraction failed: {ex}")
    try:
        raw = lollms_client.generate_text(prompt=prompt, temperature=0.1)
        json_match = re.search(r"\{.*\}", str(raw), re.DOTALL)
        if json_match:
            data = json.loads(json_match.group(0))
            if isinstance(data, dict):
                for key in result:
                    value = data.get(key)
                    if isinstance(value, str) and value.strip():
                        result[key] = value.strip()
    except Exception as ex:
        ASCIIColors.warning(f"[DocIngestor] LLM metadata extraction failed: {ex}")
    return result


class DocIngestor:
    """
    Deterministic-first document ingestion: converts txt/md/pdf/docx/pptx into a
    heading-chunked documentation tree under handbag/docs/.

    The LLM is only consulted when use_llm=True AND the deterministic heuristics
    failed to produce structure or metadata.
    """

    SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx", ".pptx"}
    _TYPE_LABELS = {
        ".md": "markdown",
        ".txt": "text",
        ".pdf": "pdf",
        ".docx": "docx",
        ".pptx": "pptx",
    }

    def ingest_document(
        self,
        file_path: Union[str, Path],
        docs_dir: Union[str, Path],
        use_llm: bool = False,
        lollms_client: Any = None,
    ) -> Path:
        source = Path(file_path)
        if not source.exists():
            raise FileNotFoundError(f"Document not found: {source}")
        ext = source.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported document format '{ext}'. Supported: {sorted(self.SUPPORTED_EXTENSIONS)}"
            )

        if ext == ".md":
            raw = source.read_text(encoding="utf-8", errors="ignore")
            md_body, fm_hints = _split_frontmatter(raw)
            hints: Dict[str, str] = {}
            if fm_hints.get("title"):
                hints["title"] = fm_hints["title"]
            author_value = fm_hints.get("author") or fm_hints.get("authors")
            if author_value:
                hints["authors"] = author_value
        elif ext == ".txt":
            md_body, hints = _convert_txt(source.read_text(encoding="utf-8", errors="ignore"))
        elif ext == ".pdf":
            md_body, hints = _convert_pdf(source)
        elif ext == ".docx":
            md_body, hints = _convert_docx(source)
        else:
            md_body, hints = _convert_pptx(source)

        title = _extract_title(md_body, hints, source.stem)
        authors = _extract_authors(md_body, hints)
        abstract = _extract_abstract(md_body)
        year = _extract_year(md_body)
        flat, preamble = _parse_markdown_sections(md_body)

        metadata_is_weak = (not flat) or (authors == "Unknown" and not abstract)
        if use_llm and lollms_client is not None and metadata_is_weak:
            llm_meta = _llm_extract_metadata(lollms_client, md_body, source.stem)
            if not hints.get("title") and llm_meta.get("title"):
                title = llm_meta["title"]
            if authors == "Unknown" and llm_meta.get("authors"):
                authors = llm_meta["authors"]
            if not abstract and llm_meta.get("abstract"):
                abstract = llm_meta["abstract"]

        docs_root = Path(docs_dir)
        docs_root.mkdir(parents=True, exist_ok=True)
        doc_dir = _unique_dir(docs_root, _safe_slug(title))
        doc_dir.mkdir(parents=True, exist_ok=True)

        (doc_dir / "full.md").write_text(md_body.strip() + "\n", encoding="utf-8")

        if not flat:
            sections: List[_Section] = [
                _Section(title=title, level=1, content=md_body.strip())
            ]
        else:
            sections = _build_section_tree(flat)
            if preamble and len(preamble) > 200:
                overview = _Section(
                    title="Overview",
                    level=sections[0].level if sections else 1,
                    content=preamble,
                )
                sections.insert(0, overview)

        root_entries = _emit_section_tree(sections, doc_dir)
        total_tokens = _estimate_tokens(md_body)

        index_fields: Dict[str, Any] = {
            "title": title,
            "authors": authors,
            "type": self._TYPE_LABELS[ext],
            "year": year,
            "source_file": source.name,
            "ingested_at": datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
        }

        lines = [_frontmatter_block(index_fields), "", f"# {title}", ""]
        year_suffix = f" ({year})" if year else ""
        lines.append(f"*Authors: {authors}{year_suffix}*")
        lines.append("")
        lines.append("## Abstract")
        lines.append(abstract if abstract else "(No abstract available)")
        lines.append("")
        lines.append("## Section Map")
        for entry in root_entries:
            if entry["is_dir"]:
                lines.append(
                    f"- [{entry['title']}]({entry['name']}/) (~{entry['tokens']} tokens, {entry['children']} subsection(s))"
                )
            else:
                lines.append(f"- [{entry['title']}]({entry['name']}) (~{entry['tokens']} tokens)")
        lines.append("")
        lines.append(f"Total estimated tokens: ~{total_tokens}")
        lines.append("")
        (doc_dir / "INDEX.md").write_text("\n".join(lines), encoding="utf-8")

        ASCIIColors.success(
            f"[DocIngestor] Ingested '{source.name}' → {doc_dir.name} "
            f"({len(root_entries)} top-level section(s), ~{total_tokens} tokens)"
        )
        return doc_dir

    def add_document(
        self,
        file_path: Union[str, Path],
        handbag_path: Union[str, Path],
        use_llm: bool = False,
        lollms_client: Any = None,
    ) -> Path:
        docs_dir = Path(handbag_path) / "docs"
        return self.ingest_document(
            file_path, docs_dir, use_llm=use_llm, lollms_client=lollms_client
        )


class DocNavigator:
    """
    Read-only navigator over an ingested handbag/docs/ tree.

    Guarantees:
    - Every returned block carries a `=== SOURCE: docs/... ===` banner.
    - Search results are grouped by document, never interleaved across sources.
    - Over-budget loads degrade to the target's index + per-child estimates.
    """

    def __init__(self, docs_dir: Union[str, Path]):
        self.docs_dir = Path(docs_dir)

    def has_documents(self) -> bool:
        if not self.docs_dir.is_dir():
            return False
        try:
            return any(item.name != ".gitkeep" for item in self.docs_dir.iterdir())
        except Exception:
            return False

    # ── document-level introspection ────────────────────────────────────────

    def _document_meta(self, doc_dir: Path) -> Dict[str, str]:
        index_path = doc_dir / "INDEX.md"
        if index_path.exists():
            try:
                _, fm = _split_frontmatter(
                    index_path.read_text(encoding="utf-8", errors="ignore")
                )
                return fm
            except Exception:
                return {}
        return {}

    def _document_tokens(self, doc_dir: Path) -> int:
        total = 0
        try:
            for file_path in doc_dir.rglob("*.md"):
                total += _estimate_tokens(
                    file_path.read_text(encoding="utf-8", errors="ignore")
                )
        except Exception:
            pass
        return total

    def _document_section_count(self, doc_dir: Path) -> int:
        try:
            return sum(1 for f in doc_dir.rglob("*.md") if f.name != "INDEX.md")
        except Exception:
            return 0

    def _document_abstract(self, doc_dir: Path) -> str:
        index_path = doc_dir / "INDEX.md"
        if not index_path.exists():
            return ""
        try:
            text = index_path.read_text(encoding="utf-8", errors="ignore")
            match = re.search(r"##\s*Abstract\s*\n+(.*?)(?=\n##\s|\Z)", text, re.DOTALL)
            if match:
                abstract = match.group(1).strip()
                if abstract and "(no abstract available)" not in abstract.lower():
                    first_sentence = re.split(r"(?<=[.!?])\s", abstract)[0]
                    return first_sentence[:200]
        except Exception:
            pass
        return ""

    def _make_doc_entry(self, doc_dir: Path, rel: str) -> Dict[str, Any]:
        meta = self._document_meta(doc_dir)
        return {
            "rel": rel,
            "name": doc_dir.name,
            "title": meta.get("title", doc_dir.name.replace("_", " ").title()),
            "authors": meta.get("authors", "Unknown"),
            "year": meta.get("year", ""),
            "type": meta.get("type", ""),
            "tokens": self._document_tokens(doc_dir),
            "sections": self._document_section_count(doc_dir),
            "abstract_digest": self._document_abstract(doc_dir),
        }

    def list_documents(self) -> List[Dict[str, Any]]:
        documents: List[Dict[str, Any]] = []
        if not self.docs_dir.is_dir():
            return documents
        try:
            top_items = sorted(self.docs_dir.iterdir(), key=lambda p: p.name)
        except Exception:
            return documents
        for top in top_items:
            if not top.is_dir() or top.name.startswith((".", "_")):
                continue
            if (top / "INDEX.md").exists():
                documents.append(self._make_doc_entry(top, rel=top.name))
            else:
                for child in sorted(top.iterdir(), key=lambda p: p.name):
                    if child.is_dir() and (child / "INDEX.md").exists():
                        documents.append(
                            self._make_doc_entry(child, rel=f"{top.name}/{child.name}")
                        )
        return documents

    # ── path resolution ─────────────────────────────────────────────────────

    def resolve(self, path: str) -> Optional[Path]:
        if not self.docs_dir.is_dir():
            return None
        cleaned = (path or "").strip().replace("\\", "/").strip("/")
        if not cleaned:
            return self.docs_dir
        lowered = cleaned.lower()
        for prefix in ("docs/", "./docs/"):
            if lowered.startswith(prefix):
                cleaned = cleaned[len(prefix):]
                break
        cleaned = cleaned.strip("/")
        if not cleaned:
            return self.docs_dir
        candidate = self.docs_dir / cleaned
        if candidate.exists():
            return candidate

        lowered = cleaned.lower()
        exact_matches: List[Path] = []
        suffix_matches: List[Path] = []
        try:
            items = list(self.docs_dir.rglob("*"))
        except Exception:
            return None
        for item in items:
            try:
                rel = item.relative_to(self.docs_dir).as_posix()
            except Exception:
                continue
            rel_lower = rel.lower()
            if rel_lower == lowered:
                exact_matches.append(item)
            elif rel_lower.endswith("/" + lowered):
                suffix_matches.append(item)
        if len(exact_matches) == 1:
            return exact_matches[0]
        if len(exact_matches) > 1:
            return None
        if len(suffix_matches) == 1:
            return suffix_matches[0]
        return None

    def _document_dir_of(self, resolved: Path) -> Optional[Path]:
        try:
            parts = resolved.relative_to(self.docs_dir).parts
        except Exception:
            return None
        if not parts:
            return None
        first = self.docs_dir / parts[0]
        if (first / "INDEX.md").exists():
            return first
        if len(parts) >= 2:
            second = first / parts[1]
            if (second / "INDEX.md").exists():
                return second
        return first if first.is_dir() else None

    def _source_banner(self, resolved: Path) -> str:
        rel = resolved.relative_to(self.docs_dir).as_posix()
        doc_dir = self._document_dir_of(resolved)
        title, authors = rel, "Unknown"
        if doc_dir is not None:
            meta = self._document_meta(doc_dir)
            title = meta.get("title") or doc_dir.name
            authors = meta.get("authors") or "Unknown"
        return f'=== SOURCE: docs/{rel} (Document: "{title}" — {authors}) ==='

    # ── navigation operations ───────────────────────────────────────────────

    def _auto_index(self, directory: Path) -> str:
        lines = [f"Index of {directory.name} (auto-generated):", ""]
        try:
            children = sorted(directory.iterdir(), key=lambda p: p.name)
        except Exception:
            return "(empty)"
        found = False
        for child in children:
            if child.name.startswith("."):
                continue
            if child.is_dir():
                lines.append(f"- {child.name}/ (~{self._document_tokens(child)} tokens)")
                found = True
            elif child.suffix.lower() == ".md":
                try:
                    tokens = _estimate_tokens(
                        child.read_text(encoding="utf-8", errors="ignore")
                    )
                except Exception:
                    tokens = 0
                lines.append(f"- {child.name} (~{tokens} tokens)")
                found = True
        return "\n".join(lines) if found else "(empty)"

    def get_index(self, path: str = "") -> str:
        resolved = self.resolve(path)
        if resolved is None:
            return f"No documentation entry matches '{path}'. Call tool_doc_index(\"\") to list the library."
        if resolved == self.docs_dir:
            documents = self.list_documents()
            if not documents:
                return "The documentation library is empty."
            lines = [f"Documentation library — {len(documents)} document(s):", ""]
            for doc in documents:
                year = f" ({doc['year']})" if doc["year"] else ""
                lines.append(
                    f"- {doc['rel']}/ — \"{doc['title']}\" — {doc['authors']}{year} "
                    f"(~{doc['tokens']:,} tokens, {doc['sections']} sections)"
                )
                if doc["abstract_digest"]:
                    lines.append(f"  Abstract digest: {doc['abstract_digest']}")
            lines.append("")
            lines.append("Call tool_doc_index(path) on any document for its full section map.")
            return "\n".join(lines)
        if resolved.is_file():
            return self.peek(path)
        banner = self._source_banner(resolved)
        index_path = resolved / "INDEX.md"
        if index_path.exists():
            content = index_path.read_text(encoding="utf-8", errors="ignore")
            if len(content) > 6000:
                content = content[:6000] + "\n... [index truncated — load sub-sections directly]"
            return f"{banner}\n\n{content}"
        return f"{banner}\n\n{self._auto_index(resolved)}"

    def _read_dir(self, directory: Path) -> str:
        parts: List[str] = []
        try:
            children = sorted(directory.iterdir(), key=lambda p: p.name)
        except Exception:
            return ""
        for child in children:
            if child.name.startswith(".") or child.name == "INDEX.md":
                continue
            if child.is_dir():
                sub = self._read_dir(child)
                if sub:
                    parts.append(sub)
            elif child.suffix.lower() == ".md":
                try:
                    parts.append(child.read_text(encoding="utf-8", errors="ignore").strip())
                except Exception:
                    pass
        return "\n\n".join(p for p in parts if p)

    def load(self, path: str, max_chars: int = _DEFAULT_MAX_CHARS) -> str:
        resolved = self.resolve(path)
        if resolved is None:
            return f"No documentation entry matches '{path}'. Call tool_doc_index(\"\") to list the library."
        if resolved == self.docs_dir:
            return self.get_index("")
        banner = self._source_banner(resolved)
        if resolved.is_dir():
            content = self._read_dir(resolved)
        else:
            try:
                content = resolved.read_text(encoding="utf-8", errors="ignore")
            except Exception as ex:
                return f"{banner}\n[Read error: {ex}]"
        content = content.strip()
        if len(content) > max_chars:
            guidance = (
                self._auto_index(resolved)
                if resolved.is_dir()
                else "(Use tool_doc_peek with a heading anchor or line number to window into this file.)"
            )
            return (
                f"{banner}\n"
                f"[TOO LARGE: ~{_estimate_tokens(content):,} tokens exceeds your current load budget "
                f"(~{max_chars // 4:,} tokens). Compose a smaller plan — index below:]\n\n{guidance}"
            )
        return f"{banner}\n[~{_estimate_tokens(content)} tokens]\n\n{content}"

    def search(self, query: str, max_results: int = 12) -> str:
        query_clean = (query or "").strip()
        if not query_clean:
            return "Empty search query."
        if not self.docs_dir.is_dir():
            return "The documentation library is empty."
        terms = [t.lower() for t in re.findall(r"\w+", query_clean) if len(t) > 2]
        query_lower = query_clean.lower()
        hits_by_doc: Dict[str, List[Tuple[int, Path, int, str]]] = {}
        for file_path in sorted(self.docs_dir.rglob("*.md")):
            try:
                file_lines = file_path.read_text(
                    encoding="utf-8", errors="ignore"
                ).splitlines()
            except Exception:
                continue
            try:
                rel_parts = file_path.relative_to(self.docs_dir).parts
            except Exception:
                continue
            if not rel_parts:
                continue
            doc_key = rel_parts[0]
            first_dir = self.docs_dir / rel_parts[0]
            if not (first_dir / "INDEX.md").exists() and len(rel_parts) >= 2:
                doc_key = f"{rel_parts[0]}/{rel_parts[1]}"
            for line_no, line in enumerate(file_lines, 1):
                line_clean = line.strip()
                if len(line_clean) < 4:
                    continue
                line_lower = line_clean.lower()
                score = 0
                if terms:
                    matched = sum(1 for term in terms if term in line_lower)
                    if matched == len(terms):
                        score = 3
                    elif matched > 0:
                        score = 2
                if score == 0 and SequenceMatcher(None, query_lower, line_lower).ratio() >= 0.65:
                    score = 1
                if score > 0:
                    hits_by_doc.setdefault(doc_key, []).append(
                        (score, file_path, line_no, line_clean)
                    )
        if not hits_by_doc:
            return f"No matches found for '{query_clean}' in the documentation library."
        sections: List[str] = []
        total = 0
        for doc_key in sorted(hits_by_doc.keys()):
            hits = sorted(hits_by_doc[doc_key], key=lambda h: h[0], reverse=True)[:3]
            doc_dir = self.docs_dir / doc_key
            meta = self._document_meta(doc_dir)
            title = meta.get("title") or doc_key
            authors = meta.get("authors") or "Unknown"
            group_lines = [
                f'=== SOURCE GROUP: docs/{doc_key}/ (Document: "{title}" — {authors}) ==='
            ]
            for _, file_path, line_no, line_clean in hits:
                rel = file_path.relative_to(self.docs_dir).as_posix()
                group_lines.append(f"- docs/{rel} (line {line_no}): {line_clean[:300]}")
                total += 1
                if total >= max_results:
                    break
            sections.append("\n".join(group_lines))
            if total >= max_results:
                break
        return (
            "\n\n".join(sections)
            + f"\n\n[{min(total, max_results)} result(s) — grouped by document. "
              "Load a full section with tool_doc_load(path) to see the complete context.]"
        )

    def peek(self, path: str, anchor: str = "", max_lines: int = _PEEK_WINDOW_LINES) -> str:
        resolved = self.resolve(path)
        if resolved is None or not resolved.is_file():
            return f"No documentation file matches '{path}'."
        banner = self._source_banner(resolved)
        try:
            lines = resolved.read_text(encoding="utf-8", errors="ignore").splitlines()
        except Exception as ex:
            return f"{banner}\n[Read error: {ex}]"
        if not lines:
            return f"{banner}\n(empty file)"
        anchor_clean = (anchor or "").strip()
        if anchor_clean.isdigit():
            center = int(anchor_clean)
            start = max(0, center - max_lines // 2)
            window = lines[start: start + max_lines]
            header = f"[lines {start + 1}-{start + len(window)} of {len(lines)}]"
            return f"{banner}\n{header}\n\n" + "\n".join(window)
        if anchor_clean:
            anchor_lower = anchor_clean.lower()
            for idx, line in enumerate(lines):
                heading_match = _HEADING_RE.match(line)
                if heading_match and anchor_lower in line.lower():
                    level = len(heading_match.group(1))
                    end = len(lines)
                    for jdx in range(idx + 1, len(lines)):
                        next_match = _HEADING_RE.match(lines[jdx])
                        if next_match and len(next_match.group(1)) <= level:
                            end = jdx
                            break
                    window = lines[idx: min(end, idx + max_lines)]
                    return (
                        f"{banner}\n[heading window: {len(window)} lines]\n\n"
                        + "\n".join(window)
                    )
        window = lines[:max_lines]
        note = (
            f"\n... [first {len(window)} of {len(lines)} lines — use an anchor (heading text or line number) to window deeper]"
            if len(lines) > max_lines
            else ""
        )
        return f"{banner}\n\n" + "\n".join(window) + note

    def build_scratchpad_block(self, path: str, note: str = "", excerpt: str = "") -> str:
        resolved = self.resolve(path)
        if resolved is None:
            return f"No documentation entry matches '{path}'."
        rel = resolved.relative_to(self.docs_dir).as_posix()
        excerpt_clean = (excerpt or "").strip()
        if excerpt_clean:
            content = self.peek(path, excerpt_clean)
        else:
            content = self.load(path)
        doc_dir = self._document_dir_of(resolved)
        doc_title = ""
        if doc_dir is not None:
            doc_title = self._document_meta(doc_dir).get("title") or doc_dir.name
        header = f"## [DOC EXTRACT] {doc_title or rel}"
        note_clean = (note or "").strip()
        note_line = f"\nNote: {note_clean}" if note_clean else ""
        return f"{header}\n{self._source_banner(resolved)}{note_line}\n\n{content}"


def _first_present(candidates: List[Any]) -> Any:
    """Returns the first non-empty candidate; tolerates positional and alias keyword arguments."""
    for candidate in candidates:
        if candidate is None:
            continue
        if str(candidate).strip():
            return candidate
    return None


def build_doc_tools(
    navigator: DocNavigator,
    scratchpad_appender: Optional[Callable[[str], str]] = None,
    max_chars_provider: Optional[Callable[[], int]] = None,
) -> Dict[str, Dict[str, Any]]:
    """Builds the 5 handbag-documentation tools. Mounted only when docs/ exists."""

    def _limit() -> int:
        if max_chars_provider is not None:
            try:
                return int(max_chars_provider())
            except Exception:
                pass
        return _DEFAULT_MAX_CHARS

    def tool_doc_index(*args, **kwargs) -> dict:
        """
        Read the index of the handbag documentation library or of one document/section.

        Args:
            path (str, optional): Document or section path (e.g. "my_document/"). Empty = library root.
        """
        path = _first_present([
            args[0] if args else None,
            kwargs.get("path"),
            kwargs.get("document_path"),
            kwargs.get("doc_path"),
        ]) or ""
        return {"success": True, "output": navigator.get_index(str(path))}

    def tool_doc_load(*args, **kwargs) -> dict:
        """
        Load the complete content of a documentation section (a file, or a folder reconstructed recursively).

        Args:
            path (str): Path of the section file or folder (e.g. "my_document/01_intro/01_motivation.md").
        """
        path = _first_present([
            args[0] if args else None,
            kwargs.get("path"),
            kwargs.get("document_path"),
            kwargs.get("doc_path"),
            kwargs.get("file_path"),
            kwargs.get("section_path"),
        ]) or ""
        if not str(path).strip():
            return {
                "success": False,
                "error": "Missing required parameter 'path'. List the library with tool_doc_index('') to discover valid paths, then call e.g. tool_doc_load('my_document/01_section.md').",
            }
        return {"success": True, "output": navigator.load(str(path), max_chars=_limit())}

    def tool_doc_search(*args, **kwargs) -> dict:
        """
        Fuzzy-search the documentation library. Results are grouped by document and always path-stamped.

        Args:
            query (str): The search keywords or question.
            max_results (int, optional): Maximum number of hits. Defaults to 12.
        """
        query = _first_present([
            args[0] if args else None,
            kwargs.get("query"),
            kwargs.get("q"),
            kwargs.get("search_query"),
            kwargs.get("search"),
            kwargs.get("keywords"),
        ]) or ""
        if not str(query).strip():
            return {
                "success": False,
                "error": "Missing required parameter 'query'. Provide the search keywords, e.g. tool_doc_search('hybrid retrieval').",
            }
        max_results = _first_present([
            kwargs.get("max_results"),
            kwargs.get("limit"),
            kwargs.get("n"),
        ]) or 12
        try:
            max_results = int(max_results)
        except (TypeError, ValueError):
            max_results = 12
        return {"success": True, "output": navigator.search(str(query), max_results=max_results)}

    def tool_doc_peek(*args, **kwargs) -> dict:
        """
        Windowed view of a large documentation file around a heading anchor or a line number.

        Args:
            path (str): Path of the documentation file.
            anchor (str, optional): A heading title fragment or a line number.
        """
        path = _first_present([
            args[0] if args else None,
            kwargs.get("path"),
            kwargs.get("document_path"),
            kwargs.get("doc_path"),
            kwargs.get("file_path"),
        ]) or ""
        if not str(path).strip():
            return {
                "success": False,
                "error": "Missing required parameter 'path'. Provide the documentation file path, e.g. tool_doc_peek('my_document/01_section.md', 'Heading Fragment').",
            }
        anchor = _first_present([
            args[1] if len(args) > 1 else None,
            kwargs.get("anchor"),
            kwargs.get("heading"),
            kwargs.get("line"),
            kwargs.get("line_number"),
        ]) or ""
        return {"success": True, "output": navigator.peek(str(path), str(anchor))}

    def tool_doc_send_to_scratchpad(*args, **kwargs) -> dict:
        """
        Send documentation content directly to the scratchpad as a source-stamped extract (no manual rewriting).

        Args:
            path (str): Path of the section or file to send.
            note (str, optional): An optional reformulated note recorded alongside the extract.
            excerpt (str, optional): Optional heading anchor or line number to send only a portion.
        """
        path = _first_present([
            args[0] if args else None,
            kwargs.get("path"),
            kwargs.get("document_path"),
            kwargs.get("doc_path"),
            kwargs.get("file_path"),
            kwargs.get("section_path"),
        ]) or ""
        if not str(path).strip():
            return {
                "success": False,
                "error": "Missing required parameter 'path'. Provide the section or file path to send, e.g. tool_doc_send_to_scratchpad('my_document/01_section.md').",
            }
        note = _first_present([
            args[1] if len(args) > 1 else None,
            kwargs.get("note"),
            kwargs.get("annotation"),
            kwargs.get("comment"),
        ]) or ""
        excerpt = _first_present([
            args[2] if len(args) > 2 else None,
            kwargs.get("excerpt"),
            kwargs.get("anchor"),
        ]) or ""
        block = navigator.build_scratchpad_block(str(path), note=str(note), excerpt=str(excerpt))
        if scratchpad_appender is None:
            return {
                "success": False,
                "error": "Scratchpad is not available (no workspace configured). "
                         "Use tool_doc_load and write a curated note instead.",
            }
        try:
            result = scratchpad_appender(block)
        except Exception as ex:
            return {"success": False, "error": f"Scratchpad append failed: {ex}"}
        if isinstance(result, str) and ("SYSTEM ERROR" in result or "❌" in result):
            return {"success": False, "error": result}
        return {
            "success": True,
            "output": f"Documentation extract sent to scratchpad.\n\n{block[:600]}",
        }

    return {
        "tool_doc_index": {
            "name": "tool_doc_index",
            "description": "Read the index of the handbag documentation library or of one document/section (metadata, abstract, section map with token estimates).",
            "parameters": [
                {"name": "path", "type": "str", "description": "Document or section path. Empty = library root.", "optional": True},
            ],
            "callable": tool_doc_index,
        },
        "tool_doc_load": {
            "name": "tool_doc_load",
            "description": "Load the complete content of a documentation section (file or folder). Over-budget loads degrade to the section index.",
            "parameters": [
                {"name": "path", "type": "str", "description": "Path of the section file or folder to load."},
            ],
            "callable": tool_doc_load,
        },
        "tool_doc_search": {
            "name": "tool_doc_search",
            "description": "Fuzzy-search the documentation library. Results are grouped by document and always carry their full source path.",
            "parameters": [
                {"name": "query", "type": "str", "description": "The search keywords or question."},
                {"name": "max_results", "type": "int", "description": "Maximum number of hits (default 12).", "optional": True},
            ],
            "callable": tool_doc_search,
        },
        "tool_doc_peek": {
            "name": "tool_doc_peek",
            "description": "Windowed view of a large documentation file around a heading anchor or a line number.",
            "parameters": [
                {"name": "path", "type": "str", "description": "Path of the documentation file."},
                {"name": "anchor", "type": "str", "description": "Heading title fragment or line number (optional).", "optional": True},
            ],
            "callable": tool_doc_peek,
        },
        "tool_doc_send_to_scratchpad": {
            "name": "tool_doc_send_to_scratchpad",
            "description": "Send documentation content directly to the scratchpad as a source-stamped extract, without rewriting it manually.",
            "parameters": [
                {"name": "path", "type": "str", "description": "Path of the section or file to send."},
                {"name": "note", "type": "str", "description": "Optional reformulated note recorded alongside the extract.", "optional": True},
                {"name": "excerpt", "type": "str", "description": "Optional heading anchor or line number to send only a portion.", "optional": True},
            ],
            "callable": tool_doc_send_to_scratchpad,
        },
    }


def build_docs_scope_block(navigator: DocNavigator, max_docs: int = _MAX_GLIMPSE_DOCS) -> str:
    """Budget-capped system-prompt glimpse of the documentation library. Empty when no docs."""
    documents = navigator.list_documents()
    if not documents:
        return ""
    total_tokens = sum(doc["tokens"] for doc in documents)
    lines = [
        "=== HANDBAG DOCUMENTATION SCOPE ===",
        f"You carry a private documentation library ({len(documents)} document(s), ~{total_tokens:,} tokens total).",
        "Documents:",
    ]
    for doc in documents[:max_docs]:
        year = f", {doc['year']}" if doc["year"] else ""
        lines.append(
            f'- "{doc["title"]}" — {doc["authors"]}{year} '
            f"(~{doc['tokens']:,} tokens, {doc['sections']} sections) → {doc['rel']}/"
        )
        if doc["abstract_digest"]:
            lines.append(f"  Abstract digest: {doc['abstract_digest']}")
    if len(documents) > max_docs:
        lines.append(
            f"... (+{len(documents) - max_docs} more — call tool_doc_index(\"\") for the full list)"
        )
    lines.append(
        "Navigate it with tool_doc_index → tool_doc_load / tool_doc_search / tool_doc_peek.\n"
        "All navigation should happen inside ONE focused documentation research phase; extract what matters into\n"
        "your scratchpad with tool_doc_send_to_scratchpad (verbatim) or curated <scratchpad_append> notes.\n"
        "Every extract must keep its [Source: docs/...] path; never merge facts from different documents into one narrative.\n"
        "External RAG knowledge bases, if any, remain separately available."
    )
    lines.append("=== END HANDBAG DOCUMENTATION SCOPE ===")
    return "\n".join(lines)